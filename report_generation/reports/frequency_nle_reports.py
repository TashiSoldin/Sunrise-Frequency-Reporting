from openpyxl import load_workbook
from openpyxl.drawing.image import Image
import pandas as pd
from enums.frequency_report_enums import LastEventTypes
from helpers.excel_helper import ExcelHelper
from helpers.os_helper import OSHelper
from helpers.datetime_helper import DatetimeHelper


class FrequencyNleReports:
    def __init__(self, data: dict, output_file_path: str) -> None:
        self.df: pd.DataFrame = data.get("content")
        self.contact_details: pd.DataFrame = data.get("contact_email")
        self.contact_email_mapping: dict = self._get_contact_email_mapping()
        self.output_file_path = output_file_path

    def _get_contact_email_mapping(self) -> dict:
        return pd.Series(
            self.contact_details["EMAIL"].values,
            index=self.contact_details["ACCNUM"],
        ).to_dict()

    def sort_df(self, df: pd.DataFrame) -> pd.DataFrame:
        df["Last Event"] = pd.Categorical(
            df["Last Event"],
            categories=LastEventTypes.get_ordered_values(),
            ordered=True,
        )
        return df.sort_values(
            by=["Last Event", "Waybill Date"], ascending=[True, False]
        )

    def generate_report(self) -> dict:
        df = self.sort_df(self.df)
        summary = {}
        template_path = OSHelper.load_template(
            "assets/", "frequency_nle_report_template.xlsx"
        )
        file_path = f"{self.output_file_path}/last_event_report-{DatetimeHelper.get_current_datetime()}.xlsx"
        logo_path = OSHelper.load_template("assets/", "logo.png")
        wb = load_workbook(template_path, data_only=False)
        template_sheet = wb["Template"]

        orighubs = (
            df["Orig Hub"].unique().tolist()
        )  # List of all the unique origin hub names

        # Loop through for each origin hub, creating a new sheet for each one
        for orighub in orighubs:
            df_orighub = df[df["Orig Hub"] == orighub]
            target_sheet = wb.copy_worksheet(
                template_sheet
            )  # Create a new sheet for the current origin hub
            target_sheet.title = str(orighub)

            # Add the logo to the new sheet since it was not copied from the template
            img = Image(logo_path)
            original_width = img.width
            original_height = img.height
            aspect_ratio = original_width / original_height
            img.width = 493
            img.height = img.width / aspect_ratio
            img.anchor = "A1"
            target_sheet.add_image(img)

            replacements = {
                "client_name": str(
                    orighub
                ),  # Client name should be the name of the origin hub
                "date_time": DatetimeHelper.get_precise_current_datetime(),
            }
            ExcelHelper.update_template_sheet(wb, orighub, replacements)

            ExcelHelper.append_df_to_sheet_with_styling(
                wb[str(orighub)],
                df_orighub,
                wb[str(orighub)].max_row + 2,  # Append the data to the new sheet
            )

        # Remove the original template sheet
        wb.remove(template_sheet)

        ExcelHelper.autofit_workbook_columns(wb)
        wb.save(file_path)

        summary["Frequency NLE report"] = {
            "file_path": file_path,
            "client_name": "All Origin Hubs",
            "email": None,
        }

        return summary
