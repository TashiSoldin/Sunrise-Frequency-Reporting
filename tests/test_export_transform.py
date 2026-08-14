"""The lifted export transform must equal a write/read round trip.

Day 3 of the DB query interface lifts the row-building loop out of write_xlsx
(export_rows / export_shaped) so a live query can feed the report builders
without producing an export file first. The contract is exact equivalence:
export_shaped(db_cols, rows) must return cell-for-cell what load_export reads
back from a file that write_xlsx wrote from the same rows — because every
report figure was reconciled against Parcel Perfect through that file path,
and the injected path inherits that verification only if the shapes match.
"""

import sys
from datetime import date, time
from decimal import Decimal
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "revenue_reports"))

from extract_revenue import (
    CREDITS_HEADERS,
    DB_COLS,
    EXPORT_HEADERS,
    credits_shaped,
    export_rows,
    export_shaped,
    extraction_sql,
    totals_row,
    write_credits_xlsx,
    write_xlsx,
)

from data import load_credit_sheet, load_export

# fetch() returns cursor-description order: the SELECT lists DB_COLS then the
# two join aliases.
FETCH_COLS = [*DB_COLS, "RECEIPT_USERNAME", "WB_TOTSURCHARGE"]


def db_row(**fields) -> tuple:
    """A raw DB row: None everywhere except the named columns."""
    ix = {c: i for i, c in enumerate(FETCH_COLS)}
    row = [None] * len(FETCH_COLS)
    for k, v in fields.items():
        row[ix[k]] = v
    return tuple(row)


ROWS = [
    db_row(
        WAYBILL="SLJNB100001",
        WAYDATE=date(2026, 8, 12),
        SERVICE="ON",
        CUSTNAME="Tommy PDY Limited",
        ACCNUM="TOM001",
        ORIGHUB="JNB",
        PIECES=5,
        ACTKG=Decimal("12.40"),
        CHARGEMASS=Decimal("20.00"),
        SUBTOTAL=Decimal("150.55"),
        STATUS="Invoiced",
        REP="TF",
        REPNAME="Tommy Fields",
        INPUTMETHOD="1",
        COLSTATUS="C",
        INSURANCEFLAG="Y",
        PODCAPTURETIME=time(9, 30, 15, 123456),
    ),
    # Consolidation child: Invoiced, zero subtotal, chargeable service,
    # external account -> Consolidated must derive "Yes".
    db_row(
        WAYBILL="SLJNB100002",
        WAYDATE=date(2026, 8, 12),
        SERVICE="ON",
        CUSTNAME="Tommy PDY Limited",
        ACCNUM="TOM001",
        ORIGHUB="JNB",
        PIECES=1,
        CHARGEMASS=Decimal("0.00"),
        SUBTOTAL=Decimal("0.00"),
        STATUS="Invoiced",
    ),
    # No-charge service at zero: NOT a consolidation. Blank customs group
    # default-fills to "Documents" on every row.
    db_row(
        WAYBILL="SLJNB100003",
        WAYDATE=date(2026, 8, 13),
        SERVICE="NCH",
        CUSTNAME="Sunrise Internal",
        ACCNUM="000",
        SUBTOTAL=Decimal("0.00"),
        CHARGEMASS=Decimal("0.00"),
        STATUS="Invoiced",
        CUSTOMSGROUP="Freight",
    ),
]


class TestDerivedFields:
    def test_r_per_kg_customs_group_and_consolidated(self):
        out = export_rows(FETCH_COLS, ROWS)
        ix = {h: i for i, h in enumerate(EXPORT_HEADERS)}
        assert out[0][ix["Avg R per kg"]] == round(150.55 / 20.00, 2)
        assert out[1][ix["Avg R per kg"]] is None  # zero mass -> blank
        assert out[0][ix["Customs Group"]] == "Documents"  # default fill
        assert out[2][ix["Customs Group"]] == "Freight"
        assert out[0][ix["Consolidated"]] == "No"
        assert out[1][ix["Consolidated"]] == "Yes"
        assert out[2][ix["Consolidated"]] == "No"  # no-charge service

    def test_value_maps_bools_and_time_truncation(self):
        out = export_rows(FETCH_COLS, ROWS)
        ix = {h: i for i, h in enumerate(EXPORT_HEADERS)}
        assert out[0][ix["Waybill Input Method"]] == "PPOnline"
        assert out[0][ix["Collect Status"]] == "Collected"
        assert out[0][ix["Insurance"]] is True
        assert out[0][ix["POD Capture Time"]] == time(9, 30, 15)


class TestRoundTrip:
    """export_shaped == write_xlsx -> load_export, cell for cell."""

    def test_shaped_rows_match_a_written_file(self, tmp_path):
        path = str(tmp_path / "roundtrip.xlsx")
        write_xlsx(path, FETCH_COLS, ROWS)
        f_headers, f_rows = load_export(path)
        s_headers, s_rows = export_shaped(FETCH_COLS, ROWS)

        assert f_headers == s_headers == EXPORT_HEADERS
        assert len(f_rows) == len(s_rows) + 1  # file carries the totals row
        for i, (f_row, s_row) in enumerate(zip(f_rows, s_rows)):
            for h, fv, sv in zip(EXPORT_HEADERS, f_row, s_row):
                assert fv == sv, f"row {i}, column {h!r}: file {fv!r} != shaped {sv!r}"

    def test_totals_row_matches_the_written_file(self, tmp_path):
        path = str(tmp_path / "totals.xlsx")
        write_xlsx(path, FETCH_COLS, ROWS)
        _, f_rows = load_export(path)
        expected = totals_row(export_rows(FETCH_COLS, ROWS))
        for h, fv, ev in zip(EXPORT_HEADERS, f_rows[-1], expected):
            ev = "" if ev is None else ev  # blanks read back as ''
            assert fv == ev, f"totals column {h!r}: file {fv!r} != computed {ev!r}"


# --- credits ----------------------------------------------------------------

CREDIT_COLS = [
    "RECEIPT",
    "ACCNUM",
    "CUSTNAME",
    "RECDATE",
    "AMOUNT",
    "DISCOUNT",
    "REFERENCE",
    "RECTYPE",
    "REASON",
    "ALLOCATED",
    "COMMENT",
    "AIF",
    "USERNAME",
    "EXPORT",
    "VAT",
    "VATTYPE",
    "BRANCHNAME",
    "BANK",
    "CUSTOMSVAT",
    "CUSTOMSDUTIES",
    "REPNAME",
    "COSTCNTRNAME",
    "CREDCONTROLLER",
]


def credit_row(**fields) -> tuple:
    ix = {c: i for i, c in enumerate(CREDIT_COLS)}
    row = [None] * len(CREDIT_COLS)
    for k, v in fields.items():
        row[ix[k]] = v
    return tuple(row)


CREDIT_ROWS = [
    credit_row(
        RECEIPT=-1001,
        ACCNUM="TOM001",
        CUSTNAME="Tommy PDY Limited",
        RECDATE=date(2026, 8, 5),
        AMOUNT=Decimal("-115.00"),
        VAT=Decimal("-15.00"),
        RECTYPE="N",
        REASON="Rate query",
        ALLOCATED=Decimal("-115.00"),
        USERNAME="MARI",
    ),
    credit_row(
        RECEIPT=-1002,
        ACCNUM="ABC002",
        CUSTNAME="ABC Freight",
        RECDATE=date(2026, 8, 6),
        AMOUNT=Decimal("-57.50"),
        VAT=Decimal("-7.50"),
        RECTYPE="J",
    ),
]


class TestCreditsRoundTrip:
    def test_shaped_rows_match_a_written_file(self, tmp_path):
        path = str(tmp_path / "credits.xlsx")
        write_credits_xlsx(path, CREDIT_COLS, CREDIT_ROWS)
        f_headers, f_rows = load_credit_sheet(path)
        s_headers, s_rows = credits_shaped(CREDIT_COLS, CREDIT_ROWS)

        assert f_headers == s_headers == CREDITS_HEADERS
        # load_credit_sheet drops the totals row (non-negative Receipt)
        assert len(f_rows) == len(s_rows)
        for i, (f_row, s_row) in enumerate(zip(f_rows, s_rows)):
            for h, fv, sv in zip(CREDITS_HEADERS, f_row, s_row):
                assert fv == sv, f"row {i}, column {h!r}: file {fv!r} != shaped {sv!r}"

    def test_rectype_maps_to_display_names(self):
        _, rows = credits_shaped(CREDIT_COLS, CREDIT_ROWS)
        ti = CREDITS_HEADERS.index("Type")
        assert [r[ti] for r in rows] == ["Credit Note", "Journal Credit"]


# --- extraction_sql parameterisation ----------------------------------------


class TestExtractionSqlFilters:
    START = date(2026, 3, 1)

    def test_account_filter_is_bound_never_interpolated(self):
        hostile = "X'; DELETE FROM WAYBILL; --"
        sql, params = extraction_sql("inv", self.START, account=hostile)
        assert hostile not in sql
        assert "wba.ACCNUM = ?" in sql
        assert params == [self.START, date(2027, 3, 1), hostile]

    def test_period_filters_bound_the_basis_column(self):
        sql, params = extraction_sql(
            "inv",
            self.START,
            date_from=date(2026, 7, 1),
            date_to=date(2026, 7, 31),
        )
        assert sql.count("wba.INVDATE >= ?") == 2  # FY floor + period floor
        assert "wba.INVDATE <= ?" in sql
        assert params == [
            self.START,
            date(2027, 3, 1),
            date(2026, 7, 1),
            date(2026, 7, 31),
        ]

    def test_wb_basis_filters_the_waybill_date(self):
        sql, params = extraction_sql("wb", self.START, date_from=date(2026, 8, 13))
        assert "wba.WAYDATE >= ?" in sql
        assert "INVDATE" not in sql.split("WHERE")[1]  # filters hit WAYDATE only
        assert params == [self.START, date(2027, 3, 1), date(2026, 8, 13)]

    def test_unknown_basis_fails_loudly(self):
        import pytest

        with pytest.raises(KeyError):
            extraction_sql("credits", self.START)  # only 'wb' and 'inv' exist

    def test_no_values_are_interpolated_into_the_sql(self):
        sql, _ = extraction_sql(
            "wb",
            self.START,
            account="TOM001",
            date_from=date(2026, 7, 1),
            date_to=date(2026, 7, 31),
        )
        assert "2026" not in sql
        assert "TOM001" not in sql
        assert "DATE '" not in sql


# --- builder injection --------------------------------------------------------


def assert_workbooks_identical(path_a: str, path_b: str):
    """Cell-for-cell equality of two workbooks, twice over: the formula view
    (literals + formula strings) and the cached-value view (data_only=True,
    what Excel-for-web/SharePoint/mobile render without recalculating). The
    second pass is what diff_workbooks.py cannot see — a formula written
    without its cached value renders 0 outside desktop Excel."""
    import openpyxl

    for data_only in (False, True):
        wa = openpyxl.load_workbook(path_a, data_only=data_only)
        wb_ = openpyxl.load_workbook(path_b, data_only=data_only)
        assert wa.sheetnames == wb_.sheetnames
        view = "cached" if data_only else "formula"
        for name in wa.sheetnames:
            sa, sb = wa[name], wb_[name]
            assert (sa.max_row, sa.max_column) == (sb.max_row, sb.max_column), (
                f"sheet {name!r}: extent differs "
                f"({sa.max_row}x{sa.max_column} vs {sb.max_row}x{sb.max_column})"
            )
            for row_a, row_b in zip(sa.iter_rows(), sb.iter_rows()):
                for ca, cb in zip(row_a, row_b):
                    assert ca.value == cb.value, (
                        f"{name}!{ca.coordinate} ({view} view): "
                        f"file-path {ca.value!r} != injected {cb.value!r}"
                    )


class TestBuilderInjection:
    """build(data=...) must produce the identical workbook to build from the
    file written from the same rows — the injected path inherits the file
    path's verification only if the outputs are indistinguishable."""

    def test_flash_from_injected_rows_equals_flash_from_file(self, tmp_path):
        from build_flash import build

        export = str(tmp_path / "WB Date - test.xlsx")
        write_xlsx(export, FETCH_COLS, ROWS)
        day = date(2026, 8, 12)

        out_file = tmp_path / "from_file"
        out_data = tmp_path / "from_data"
        out_file.mkdir()
        out_data.mkdir()
        a = build(day, export, str(out_file))
        b = build(day, None, str(out_data), data=export_shaped(FETCH_COLS, ROWS))
        assert_workbooks_identical(a, b)


# --- the four builders Day 3 left untested ------------------------------------
# Fixture sized to actually cross the guards the builders carry: >= 20 waybills
# on each day the billing frontier must call "done" (build_unbilled.MIN_WAYBILLS
# — a fixture under a threshold executes none of the rule it names), several
# accounts/reps/branches, a consolidation group with zero-charged children, a
# zero-billed group, unbilled statuses, a tilde waybill, an internal account,
# a DEDI-folded account, and prior-FY + credit files for the dashboard.

_ACCTS = [
    # acct, customer, salesrep, rep name, orig hub
    ("TOM001", "Tommy PDY Limited", "TF", "Tracy Flandorp", "JNB"),
    ("ABC002", "ABC Freight", "CN", "Christine Naidoo", "CPT"),
    ("XYZ003", "XYZ Mining", "LS", "Larry Serman", "DUR"),
    ("PLZ004", "Bay Traders", "TF", "Tracy Flandorp", "PLZ"),  # branch "Other"
]


def _way(i, wd, acct_ix, *, status="Invoiced", invdate=None, invoice=None, sub, kg):
    acct, cust, rep, repname, hub = _ACCTS[acct_ix % len(_ACCTS)]
    sub = Decimal(sub)
    return db_row(
        WAYBILL=f"SLJNB{100100 + i}",
        WAYDATE=wd,
        SERVICE="ON",
        CUSTNAME=cust,
        ACCNUM=acct,
        INVOICE=invoice
        if invoice is not None
        else (0 if invdate is None else 2000 + i),
        INVDATE=invdate,
        REFERENCE=f"PO-{7000 + i}",
        ORIGPERS=f"Shipper {i}",
        DESTPERS=f"Consignee {i}",
        ORIGHUB=hub,
        DESTHUB="JNB" if hub != "JNB" else "CPT",
        DESTTOWN="Springs",
        PIECES=1 + i % 4,
        ACTKG=Decimal("2.50") * (1 + i % 3),
        CHARGEMASS=Decimal(kg),
        SUBTOTAL=sub,
        CARTAGE=round(sub * Decimal("0.6"), 2),
        SURCHARGE6=round(sub * Decimal("0.1"), 2),
        DOCS=Decimal("5.25"),
        OUTLY=Decimal("14.80") if i % 5 == 0 else None,
        STATUS=status,
        REP=rep,
        REPNAME=repname,
    )


def _fy27_rows():
    rows = []
    i = 0
    # 10 Aug: 20 invoiced (11 Aug) + 2 unbilled -> frontier calls the day done
    for _ in range(20):
        rows.append(
            _way(
                i,
                date(2026, 8, 10),
                i,
                invdate=date(2026, 8, 11),
                sub=f"{100 + (i % 7) * 37}.35",
                kg=f"{10 + i % 9}.0",
            )
        )
        i += 1
    for status in ("Ready for Approval", "Quote"):
        rows.append(_way(i, date(2026, 8, 10), i, status=status, sub="88.20", kg="4.0"))
        i += 1
    # 11 Aug: 24 invoiced (12 Aug) + 3 unbilled + consolidations + specials
    for _ in range(24):
        rows.append(
            _way(
                i,
                date(2026, 8, 11),
                i,
                invdate=date(2026, 8, 12),
                sub=f"{150 + (i % 5) * 61}.10",
                kg=f"{8 + i % 6}.5",
            )
        )
        i += 1
    for status in ("Checked In", "Collection", "Summary Capture"):
        rows.append(_way(i, date(2026, 8, 11), i, status=status, sub="42.75", kg="3.0"))
        i += 1
    # consolidation group: master billed, two zero-charged children ("Yes")
    for sub in ("500.00", "0.00", "0.00"):
        r = _way(i, date(2026, 8, 11), 0, invdate=date(2026, 8, 12), sub=sub, kg="6.0")
        rows.append(_with(r, WAYREF_FIRST="CONS01", WAYREF_COUNT=3))
        i += 1
    # zero-billed group: nothing charged anywhere -> "NOT BILLED" band
    for _ in range(2):
        r = _way(
            i, date(2026, 8, 11), 1, invdate=date(2026, 8, 12), sub="0.00", kg="2.0"
        )
        rows.append(_with(r, WAYREF_FIRST="CONS02", WAYREF_COUNT=2))
        i += 1
    # tilde re-delivery duplicate: excluded by every builder
    r = _way(i, date(2026, 8, 11), 0, invdate=date(2026, 8, 12), sub="99.99", kg="1.0")
    rows.append(_with(r, WAYBILL="SLJNB900001~1"))
    i += 1
    # internal no-charge and a DEDI-folded dedicated account
    r = _way(i, date(2026, 8, 11), 0, invdate=date(2026, 8, 12), sub="0.00", kg="0.0")
    rows.append(_with(r, ACCNUM="000", CUSTNAME="Sunrise Internal", SERVICE="NCH"))
    i += 1
    r = _way(
        i, date(2026, 8, 11), 0, invdate=date(2026, 8, 12), sub="220.00", kg="11.0"
    )
    rows.append(_with(r, ACCNUM="TOM001DED", CUSTNAME="Tommy PDY Dedicated"))
    i += 1
    # house and non-budget (salesrep AN -> NEW) activity
    r = _way(i, date(2026, 8, 11), 2, invdate=date(2026, 8, 12), sub="75.00", kg="5.0")
    rows.append(_with(r, ACCNUM="HOU001", CUSTNAME="Sunrise House"))
    i += 1
    r = _way(i, date(2026, 8, 11), 3, invdate=date(2026, 8, 12), sub="130.00", kg="6.5")
    rows.append(
        _with(
            r,
            ACCNUM="NEW005",
            CUSTNAME="Fresh Logistics",
            REP="AN",
            REPNAME="Adrian van Niekerk",
        )
    )
    i += 1
    # 12 Aug: 6 rows, one invoiced 13 Aug -> the day stays below MIN_WAYBILLS
    rows.append(
        _way(
            i, date(2026, 8, 12), 0, invdate=date(2026, 8, 13), sub="310.40", kg="12.0"
        )
    )
    i += 1
    for status in (
        "Ready for Approval",
        "Quote",
        "Checked In",
        "No Rate Found",
        "Recalc Required",
    ):
        rows.append(_way(i, date(2026, 8, 12), i, status=status, sub="55.10", kg="2.5"))
        i += 1
    # completed FY27 months for the dashboard trend (Mar-Jul)
    for m in (3, 4, 5, 6, 7):
        for j in range(2):
            rows.append(
                _way(
                    i,
                    date(2026, m, 10 + j),
                    j,
                    invdate=date(2026, m, 12 + j),
                    sub=f"{400 + m * 13}.55",
                    kg="20.0",
                )
            )
            i += 1
    return rows


def _with(row: tuple, **overrides) -> tuple:
    ix = {c: j for j, c in enumerate(FETCH_COLS)}
    out = list(row)
    for k, v in overrides.items():
        out[ix[k]] = v
    return tuple(out)


def _fy26_rows():
    rows, i = [], 500
    for m in (3, 5, 8):
        for j in range(2):
            rows.append(
                _way(
                    i,
                    date(2025, m, 9 + j),
                    j,
                    invdate=date(2025, m, 11 + j),
                    sub=f"{350 + m * 11}.45",
                    kg="18.0",
                )
            )
            i += 1
    # a closed account that only traded last year
    r = _way(i, date(2025, 8, 12), 0, invdate=date(2025, 8, 13), sub="60.00", kg="3.0")
    rows.append(_with(r, ACCNUM="OLD001", CUSTNAME="Bygone Freight"))
    return rows


_CREDITS = [
    credit_row(
        RECEIPT=-2001,
        ACCNUM="TOM001",
        CUSTNAME="Tommy PDY Limited",
        RECDATE=date(2026, 3, 15),
        AMOUNT=Decimal("-230.00"),
        VAT=Decimal("-30.00"),
        RECTYPE="N",
        REASON="Rate query",
        ALLOCATED=Decimal("-230.00"),
        USERNAME="MARI",
        REPNAME="Tracy Flandorp",
        BRANCHNAME="JHB",
        CREDCONTROLLER="Mari V",
        REFERENCE="CN-1",
    ),
    credit_row(
        RECEIPT=-2002,
        ACCNUM="ABC002",
        CUSTNAME="ABC Freight",
        RECDATE=date(2026, 5, 20),
        AMOUNT=Decimal("-115.00"),
        VAT=Decimal("-15.00"),
        RECTYPE="N",
        REASON="Damaged goods",
        USERNAME="MARI",
        REPNAME="Christine Naidoo",
        BRANCHNAME="CPT",
        CREDCONTROLLER="Mari V",
    ),
    credit_row(
        RECEIPT=-2003,
        ACCNUM="TOM001",
        CUSTNAME="Tommy PDY Limited",
        RECDATE=date(2026, 8, 5),
        AMOUNT=Decimal("-345.00"),
        VAT=Decimal("-45.00"),
        RECTYPE="N",
        REASON="Rate query",
        ALLOCATED=Decimal("-115.00"),
        USERNAME="MARI",
        REPNAME="Tracy Flandorp",
        BRANCHNAME="JHB",
        CREDCONTROLLER="Mari V",
    ),
    credit_row(
        RECEIPT=-2004,
        ACCNUM="ABC002",
        CUSTNAME="ABC Freight",
        RECDATE=date(2026, 8, 12),
        AMOUNT=Decimal("-57.50"),
        VAT=Decimal("-7.50"),
        RECTYPE="N",
        REASON="Damaged goods",
        USERNAME="MARI",
        REPNAME="Christine Naidoo",
        BRANCHNAME="CPT",
        CREDCONTROLLER="Mari V",
    ),
    credit_row(
        RECEIPT=-2005,
        ACCNUM="XYZ003",
        CUSTNAME="XYZ Mining",
        RECDATE=date(2026, 8, 12),
        AMOUNT=Decimal("-23.00"),
        VAT=Decimal("-3.00"),
        RECTYPE="J",
        REASON="Journal",
        USERNAME="MARI",
        REPNAME="Larry Serman",
        BRANCHNAME="DUR",
        CREDCONTROLLER="Mari V",
    ),
    # excluded types: bad debt and cancelled must not reach any report
    credit_row(
        RECEIPT=-2006,
        ACCNUM="TOM001",
        CUSTNAME="Tommy PDY Limited",
        RECDATE=date(2026, 8, 6),
        AMOUNT=Decimal("-999.00"),
        RECTYPE="B",
    ),
    credit_row(
        RECEIPT=-2007,
        ACCNUM="TOM001",
        CUSTNAME="Tommy PDY Limited",
        RECDATE=date(2026, 8, 7),
        AMOUNT=Decimal("-111.00"),
        RECTYPE="X",
    ),
    # prior-year credit for the dashboard's LY nets
    credit_row(
        RECEIPT=-1050,
        ACCNUM="TOM001",
        CUSTNAME="Tommy PDY Limited",
        RECDATE=date(2025, 8, 5),
        AMOUNT=Decimal("-100.00"),
        RECTYPE="N",
        REASON="Rate query",
        USERNAME="MARI",
        REPNAME="Tracy Flandorp",
        BRANCHNAME="JHB",
        CREDCONTROLLER="Mari V",
    ),
]


def _budget_file(tmp_path) -> str:
    """A minimal Budget v30 lookalike: sheet 'Client Monthly Compare', headers
    on the fifth row with 'Acct' at index 3 and 'Mar...' at index 5, group
    header rows (text in col 0, blank Acct), account rows with branch/acct/name
    at 2/3/4 and monthly targets at columns 6, 8, ... 28 (Mar..Feb)."""
    from openpyxl import Workbook

    def acct(branch, a, name, monthly):
        row = [""] * 30
        row[2], row[3], row[4] = branch, a, name
        for k in range(12):
            row[6 + 2 * k] = monthly
        return row

    def group(label):
        row = [""] * 30
        row[0] = label
        return row

    wb = Workbook()
    ws = wb.active
    ws.title = "Client Monthly Compare"
    for _ in range(4):
        ws.append(["FY26-27 Budget"] + [""] * 29)
    hdr = [""] * 30
    hdr[3], hdr[4], hdr[5] = "Acct", "Client", "Mar 26"
    ws.append(hdr)
    for row in [
        group("TF — Tracy Flandorp"),
        acct("JHB", "TOM001", "Tommy PDY Limited", 50000),
        acct("JHB", "PLZ004", "Bay Traders", 10000),
        group("CN — Christine Naidoo"),
        acct("CPT", "ABC002", "ABC Freight", 30000),
        group("LS — Larry Serman"),
        acct("DUR", "XYZ003", "XYZ Mining", 20000),
        group("DEDI"),
        acct("JHB", "TOM001DED", "Tommy PDY Dedicated", 0),
        group("HS"),
        acct("JHB", "HOU001", "Sunrise House", 0),
        group("CLOSED"),
        acct("JHB", "OLD001", "Bygone Freight", 0),
    ]:
        ws.append(row)
    path = str(tmp_path / "FY26-27_Budget_v30_test.xlsx")
    wb.save(path)
    return path


def _dirs(tmp_path):
    a, b = tmp_path / "from_file", tmp_path / "from_data"
    a.mkdir()
    b.mkdir()
    return str(a), str(b)


class TestUnbilledInjection:
    def test_unbilled_from_injected_rows_equals_file(self, tmp_path):
        from build_unbilled import build

        rows = _fy27_rows()
        export = str(tmp_path / "WB Date - test.xlsx")
        write_xlsx(export, FETCH_COLS, rows)
        out_a, out_b = _dirs(tmp_path)
        a = build(export, out_a)
        b = build(export, out_b, data=export_shaped(FETCH_COLS, rows))
        assert_workbooks_identical(a, b)

    def test_unbilled_detail_is_populated(self, tmp_path):
        """The frontier fixture must exercise the rule, not duck under it:
        5 unbilled waybills dated before the 12 Aug frontier."""
        import openpyxl
        from build_unbilled import build

        out = build(
            "unused.xlsx", str(tmp_path), data=export_shaped(FETCH_COLS, _fy27_rows())
        )
        ws = openpyxl.load_workbook(out)["Unbilled Detail"]
        detail = [
            r for r in ws.iter_rows(min_row=2) if r[1].value and r[0].value != "TOTAL"
        ]
        assert len(detail) == 5

    def test_unbilled_accepts_no_source_file(self, tmp_path):
        """A live-query caller has no file to name; the Overview note must not
        crash on wb_file=None (it is only a label on this path)."""
        import openpyxl
        from build_unbilled import build

        out = build(None, str(tmp_path), data=export_shaped(FETCH_COLS, _fy27_rows()))
        ws = openpyxl.load_workbook(out)["Overview"]
        assert "live query" in ws["A3"].value


class TestBillingDetailInjection:
    def test_billing_detail_from_injected_rows_equals_file(self, tmp_path):
        from build_billing_detail import build

        rows = _fy27_rows()
        inv = str(tmp_path / "INV Date - test.xlsx")
        write_xlsx(inv, FETCH_COLS, rows)
        credits = str(tmp_path / "Credits - test.xlsx")
        write_credits_xlsx(credits, CREDIT_COLS, _CREDITS)
        out_a, out_b = _dirs(tmp_path)
        day = date(2026, 8, 12)
        a = build(day, inv, credits, out_a)
        b = build(
            day,
            inv,
            credits,
            out_b,
            inv_data=export_shaped(FETCH_COLS, rows),
            credits_data=credits_shaped(CREDIT_COLS, _CREDITS),
        )
        assert_workbooks_identical(a, b)

    def test_fixture_reaches_all_tabs(self, tmp_path):
        """Consolidations must contain both a billed-on-master group and a
        NOT-BILLED group; the credit-notes tab must carry the day's two notes."""
        import openpyxl
        from build_billing_detail import build

        out = build(
            date(2026, 8, 12),
            "unused",
            "unused",
            str(tmp_path),
            inv_data=export_shaped(FETCH_COLS, _fy27_rows()),
            credits_data=credits_shaped(CREDIT_COLS, _CREDITS),
        )
        wb_ = openpyxl.load_workbook(out)
        assert wb_.sheetnames == [
            "Billing 12 Aug",
            "By Customer",
            "Consolidations",
            "Credit Notes 12 Aug",
        ]
        cons = [
            c.value for row in wb_["Consolidations"].iter_rows() for c in row if c.value
        ]
        assert any("NOT BILLED" in str(v) for v in cons)
        assert any("billed on master" in str(v) for v in cons)
        cn = [c.value for row in wb_["Credit Notes 12 Aug"].iter_rows() for c in row]
        assert "-2004" in cn and "-2005" in cn  # the day's Credit Note + Journal
        assert "-2006" not in cn and "-2007" not in cn  # Bad Debt / Cancelled


class TestCreditNotesInjection:
    def test_credit_notes_from_injected_rows_equals_file(self, tmp_path):
        from build_credit_notes import build

        credits = str(tmp_path / "Credits - test.xlsx")
        write_credits_xlsx(credits, CREDIT_COLS, _CREDITS)
        out_a, out_b = _dirs(tmp_path)
        a = build(credits, out_a)  # month=None: derived from the data
        b = build(credits, out_b, data=credits_shaped(CREDIT_COLS, _CREDITS))
        assert_workbooks_identical(a, b)


class TestDashboardInjection:
    def test_dashboard_from_injected_rows_equals_file(self, tmp_path):
        from build_dashboard import build

        fy27, fy26 = _fy27_rows(), _fy26_rows()
        inv = str(tmp_path / "INV Date - test.xlsx")
        py_inv = str(tmp_path / "INV Date - prior.xlsx")
        wb_file = str(tmp_path / "WB Date - test.xlsx")
        credits = str(tmp_path / "Credits - test.xlsx")
        write_xlsx(inv, FETCH_COLS, fy27)
        write_xlsx(py_inv, FETCH_COLS, fy26)
        write_xlsx(wb_file, FETCH_COLS, fy27)
        write_credits_xlsx(credits, CREDIT_COLS, _CREDITS)
        budget = _budget_file(tmp_path)
        out_a, out_b = _dirs(tmp_path)

        a = build(inv, py_inv, wb_file, credits, budget, out_a)
        b = build(
            inv,
            py_inv,
            wb_file,
            credits,
            budget,
            out_b,
            inv_data=export_shaped(FETCH_COLS, fy27),
            py_inv_data=export_shaped(FETCH_COLS, fy26),
            wb_data=export_shaped(FETCH_COLS, fy27),
            credits_data=credits_shaped(CREDIT_COLS, _CREDITS),
        )
        assert_workbooks_identical(a, b)

    def test_fixture_builds_all_eleven_tabs(self, tmp_path):
        import openpyxl
        from build_dashboard import build

        fy27, fy26 = _fy27_rows(), _fy26_rows()
        out = build(
            "u",
            "u",
            "u",
            "u",
            _budget_file(tmp_path),
            str(tmp_path),
            inv_data=export_shaped(FETCH_COLS, fy27),
            py_inv_data=export_shaped(FETCH_COLS, fy26),
            wb_data=export_shaped(FETCH_COLS, fy27),
            credits_data=credits_shaped(CREDIT_COLS, _CREDITS),
        )
        names = openpyxl.load_workbook(out).sheetnames
        assert names == [
            "YTD Revenue vs PY",
            "Mar FY27",
            "Apr FY27",
            "May FY27",
            "Jun FY27",
            "Jul FY27",
            "YTD to Jul FY27",
            "MTD Billing by Customer",
            "Daily — Invoice date",
            "Daily — Waybill date",
            "MTD Credit Notes",
        ]


# --- known transform divergences, pinned with their bounds ---------------------


class TestKnownTransformDivergences:
    """Two places where export_shaped legitimately differs from the file round
    trip, found and bounded against the live DB on 14 Aug 2026. Pinned so a
    library upgrade that changes either is noticed."""

    def test_epoch_zero_date_degrades_through_the_file_path(self, tmp_path):
        """Parcel Perfect's null-date placeholder 1899-12-30 (Excel serial 0)
        reads back from a written file as time(0, 0); the injected path keeps
        the date. 6 cells in 94,431 FY27 rows (Due Date / Last Delivery Date —
        columns no builder reads). Injected is the truer value; anything that
        consumes these columns must treat pre-1990 dates as junk either way."""
        epoch = date(1899, 12, 30)
        rows = [db_row(WAYBILL="SLJNB1", WAYDATE=date(2026, 8, 12), DUEDATE=epoch)]
        path = str(tmp_path / "epoch.xlsx")
        write_xlsx(path, FETCH_COLS, rows)
        _, f_rows = load_export(path)
        _, s_rows = export_shaped(FETCH_COLS, rows)
        i = EXPORT_HEADERS.index("Due Date")
        assert s_rows[0][i] == epoch
        assert f_rows[0][i] == time(0, 0)  # the file path's degraded reading

    def test_double_precision_narrows_through_the_file_path(self, tmp_path):
        """Floats needing 17 significant digits (PP stores float32 artifacts
        like 236.92999267578125) lose their last bit through the written file;
        the injected path keeps full precision. Bounded live on the full FY27
        extract, 14 Aug 2026: 99,760 differing cells, max diff 3.6e-12, zero
        cells rounding differently at 2dp, totals row identical. Downstream it
        surfaces only in the billing detail's DERIVED Surcharge column (sub
        minus the four components, written unrounded): 8 of 1,828 cells on the
        13 Aug live build, max 6e-13 apart, identical at the displayed 2dp.
        The builder-injection tests above use 2dp fixture values, which round
        trip exactly, so they assert full equality."""
        v = 236.92999267578125
        rows = [db_row(WAYBILL="SLJNB1", WAYDATE=date(2026, 8, 12), HANDLING=v)]
        path = str(tmp_path / "ulp.xlsx")
        write_xlsx(path, FETCH_COLS, rows)
        _, f_rows = load_export(path)
        _, s_rows = export_shaped(FETCH_COLS, rows)
        i = EXPORT_HEADERS.index("Handling")
        assert s_rows[0][i] == v  # injected: exact
        assert f_rows[0][i] != v  # file: last ULP gone...
        assert abs(f_rows[0][i] - v) < 1e-9  # ...but bounded
        assert round(f_rows[0][i], 2) == round(v, 2)  # and never a cent
