"""The Billing tab inside the Flash Revenue workbook (Reuven, 19 Oct 2026).

Reuven asked for tab 1 of the Billing Detail — "Billing DD Mon", every invoice
line for the billing day — to appear in the flash as well. The two reports sit
on different date bases (flash = waybill date, billing = invoice date, one day
further in arrears), so the tab keeps its own date in its name and "Invoice
date" in its title; that is what makes the two readable side by side.

Three contracts are pinned here:

  * Sheet order. build_billing_detail._flash_comparison opens the frozen flash
    with openpyxl and reads `.active` — the FIRST sheet — scanning for three
    literal section headers. The billing tab must come after it, and the
    "All Customers" tab after that (Akha's default, 20 Oct).
  * One writer, not a copy. The tab is written by the function the Billing
    Detail itself calls, so the two workbooks carry identical cells and
    identical formatting for the same invoice lines. Asserted cell by cell.
  * The flash never fails because of the billing tab. A billing day with no
    invoice lines (month-start quiet days, a fresh deployment) still gets the
    tab — title block plus a one-line note — and an unreadable invoice source
    is skipped with a warning rather than aborting the 07:00 run.
"""

import sys
from argparse import Namespace
from datetime import date
from pathlib import Path

import openpyxl
import pytest

REPO = Path(__file__).resolve().parent.parent
BUILDERS = REPO / "revenue_reports"
sys.path.insert(0, str(BUILDERS))

import build_billing_detail
import build_flash

# Copied, not imported, for the same reason test_flash_all_customers copies
# them: they are string literals in the other builder, and the test should
# fail if the two drift apart.
PARSER_SECTIONS = ("BY BRANCH (origin)", "BY REP", "TOP 10 CUSTOMERS")

FLASH_DAY = date(2026, 8, 3)  # Monday — waybill basis
PRIOR = date(2026, 7, 27)  # same weekday, for the typical-Mon benchmark
BILLING_DAY = date(2026, 7, 31)  # Friday — invoice basis, further in arrears

WB_HEADERS = [
    "Waybill",
    "Waybill Date",
    "Customer",
    "Account",
    "Orig Hub",
    "Chrg Mass",
    "Subtotal",
    "Salesrep",
    "Rep",
]

INV_HEADERS = [
    "Waybill",
    "Waybill Date",
    "Invoice",
    "Account",
    "Customer",
    "Shipper",
    "Consignee",
    "Orig Hub",
    "Dest Hub",
    "Dest Place",
    "Reference",
    "Service",
    "Pieces",
    "Actual Mass",
    "Chrg Mass",
    "Basic Charge",
    "Outlying Charge",
    "Doc Charge",
    "Fuel",
    "Subtotal",
    "Invoice Date",
    "Salesrep",
    "Rep",
    "First Ref",
    "Consolidated",
]

CREDIT_HEADERS = [
    "Receipt",
    "Account",
    "Customer Name",
    "Date",
    "Subtotal",
    "Reference",
    "Type",
    "Rep",
    "Reason",
    "Branch",
    "Credit Controller",
]


def _wb_rows():
    rows = []
    for i in range(12):
        rows.append(
            [
                f"SL{i:05d}",
                FLASH_DAY,
                f"CUSTOMER {i}",
                f"A{i}",
                "JNB",
                100 + i,
                1000 - i * 10,
                "TF",
                "Tracy Flandorp",
            ]
        )
    rows.append(
        [
            "SL00100",
            FLASH_DAY,
            "CUSTOMER 0 (PTY) LTD",
            "A0",
            "CPT",
            50,
            500,
            "CN",
            "Christine Naidoo",
        ]
    )
    rows.append(
        ["SL00200", PRIOR, "CUSTOMER 1", "A1", "DUR", 10, 900, "TF", "Tracy Flandorp"]
    )
    return rows


def _inv_line(wb, inv, acc, cust, kg, basic, fuel, sub, day=BILLING_DAY, wbday=None):
    return [
        wb,
        wbday or date(2026, 7, 29),
        float(inv),  # invoice numbers come through the export as floats
        acc,
        cust,
        f"{cust} shipper",
        "Some consignee",
        "JNB",
        "CPT",
        "Cape Town",
        f"REF-{wb}",
        "ONX",
        2,
        kg - 0.5,
        kg,
        basic,
        0,
        15,
        fuel,
        sub,
        day,
        "TF",
        "Tracy Flandorp",
        "",
        "No",
    ]


def _inv_rows():
    """Two accounts on the billing day, one line on another day (excluded) and
    one tilde waybill (excluded) — the same filters load_day applies for the
    Billing Detail."""
    return [
        _inv_line("SL00001", 1013601, "A1", "CUSTOMER 1", 120.0, 800, 120, 1000),
        _inv_line("SL00002", 1013601, "A1", "CUSTOMER 1", 60.0, 400, 60, 500),
        _inv_line("SL00003", 1013602, "A0", "CUSTOMER 0", 20.0, 150, 20, 200),
        _inv_line(
            "SL00004",
            1013603,
            "A2",
            "CUSTOMER 2",
            10.0,
            90,
            10,
            100,
            day=date(2026, 7, 30),
        ),
        _inv_line("~SL00005", 1013604, "A1", "CUSTOMER 1", 5.0, 40, 5, 50),
    ]


@pytest.fixture
def flash(tmp_path):
    out = build_flash.build(
        FLASH_DAY,
        None,
        str(tmp_path),
        data=(WB_HEADERS, _wb_rows()),
        inv_file=None,
        billing_day=BILLING_DAY,
        inv_data=(INV_HEADERS, _inv_rows()),
    )
    return openpyxl.load_workbook(out)


@pytest.fixture
def billing_detail(tmp_path):
    out = build_billing_detail.build(
        BILLING_DAY,
        None,
        None,
        str(tmp_path / "bd"),
        None,
        inv_data=(INV_HEADERS, _inv_rows()),
        credits_data=(CREDIT_HEADERS, []),
    )
    return openpyxl.load_workbook(out)


@pytest.fixture(autouse=True)
def _bd_dir(tmp_path):
    (tmp_path / "bd").mkdir(exist_ok=True)


class TestSheetOrder:
    def test_flash_first_billing_second_all_customers_third(self, flash):
        assert flash.sheetnames == [
            "Flash 03 Aug",
            "Billing 31 Jul",
            build_flash.ALL_TAB,
        ]

    def test_the_flash_page_is_still_the_active_sheet(self, flash):
        """openpyxl's .active is what _flash_comparison opens."""
        assert flash.active.title == "Flash 03 Aug"
        assert flash.sheetnames[0] == flash.active.title

    def test_the_comparison_parser_still_sees_reps_not_invoice_lines(self, flash):
        """Replays build_billing_detail._flash_comparison against sheet one."""
        f, section, reps, branches = flash.active, None, {}, {}
        for row in f.iter_rows(min_col=2, max_col=3):
            v = row[0].value
            if v in PARSER_SECTIONS:
                section = v
            elif (
                section == "BY BRANCH (origin)"
                and v
                and row[1].value is not None
                and v != "Branch"
            ):
                branches[v] = row[1].value
            elif section == "BY REP" and v and row[1].value is not None and v != "Rep":
                reps[str(v).split(" — ")[0]] = row[1].value
        assert set(reps) == {"TF", "CN"}
        assert set(branches) == {"JHB", "Cape Town"}

    def test_without_billing_arguments_the_flash_is_unchanged(self, tmp_path):
        """The CLI contract build_flash had before this change still holds."""
        out = build_flash.build(
            FLASH_DAY, None, str(tmp_path), data=(WB_HEADERS, _wb_rows())
        )
        wb = openpyxl.load_workbook(out)
        assert wb.sheetnames == ["Flash 03 Aug", build_flash.ALL_TAB]


class TestTheTabSaysWhichDateItIs:
    def test_sheet_name_carries_the_billing_day(self, flash):
        assert "Billing 31 Jul" in flash.sheetnames

    def test_title_carries_the_invoice_date_basis(self, flash):
        ws = flash["Billing 31 Jul"]
        assert ws["B3"].value == "Billing Component Detail — Invoice date 31 Jul 2026"
        assert "Invoice date" in ws["B3"].value
        assert "invoiced on 31 Jul 2026" in ws["B4"].value

    def test_nothing_on_the_tab_claims_the_flash_day(self, flash):
        ws = flash["Billing 31 Jul"]
        text = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
        assert "3 August" not in text and "03 Aug" not in text

    def test_the_flash_page_keeps_its_own_waybill_basis(self, flash):
        assert "Waybill-date basis" in flash.active["B4"].value
        assert "3 August 2026" in flash.active["B3"].value


class TestSameWriterAsTheBillingDetail:
    """The lifted writer, called from two builders, must leave identical sheets."""

    def _pair(self, flash, billing_detail):
        return flash["Billing 31 Jul"], billing_detail["Billing 31 Jul"]

    def test_every_cell_value_matches(self, flash, billing_detail):
        a, b = self._pair(flash, billing_detail)
        coords = {
            c.coordinate for row in a.iter_rows() for c in row if c.value is not None
        } | {c.coordinate for row in b.iter_rows() for c in row if c.value is not None}
        assert coords, "empty tab"
        diffs = [
            (k, a[k].value, b[k].value)
            for k in sorted(coords)
            if a[k].value != b[k].value
        ]
        assert diffs == []

    def test_every_cell_format_matches(self, flash, billing_detail):
        """diff_workbooks.py sees values only, so the formatting is checked here."""
        a, b = self._pair(flash, billing_detail)

        def fmt(c):
            return (
                c.number_format,
                c.font.bold,
                c.font.size,
                c.font.color.rgb if c.font.color else None,
                c.fill.start_color.rgb if c.fill and c.fill.start_color else None,
                c.border.left.style,
                c.alignment.horizontal,
                c.alignment.wrap_text,
            )

        coords = {c.coordinate for row in b.iter_rows() for c in row if c.has_style}
        diffs = [k for k in sorted(coords) if fmt(a[k]) != fmt(b[k])]
        assert diffs == []

    def test_sheet_level_properties_match(self, flash, billing_detail):
        a, b = self._pair(flash, billing_detail)
        assert a.freeze_panes == b.freeze_panes == "B8"
        assert a.sheet_properties.tabColor.rgb == b.sheet_properties.tabColor.rgb
        assert a.sheet_properties.tabColor.rgb.endswith("05003C")  # TAB_SUMMARY
        for col in "ABCDEFGHIJKLMNOPQRSTUVW":
            assert a.column_dimensions[col].width == b.column_dimensions[col].width, col
        for r in range(1, 9):
            assert a.row_dimensions[r].height == b.row_dimensions[r].height, r
        assert a.row_dimensions[7].height == pytest.approx(21.9)  # H_BILLING

    def test_totals_tie_to_the_injected_lines(self, flash):
        """1000 + 500 + 200 on the day; the 30 Jul line and the tilde waybill
        are out — the same filter the Billing Detail applies."""
        ws = flash["Billing 31 Jul"]
        grand = next(
            r
            for r in range(1, ws.max_row + 1)
            if ws.cell(r, 6).value == "GRAND TOTAL — all customers"
        )
        assert ws.cell(grand, 22).value == pytest.approx(1700)
        assert ws.cell(grand, 16).value == pytest.approx(200.0)  # chargeable kg
        assert "3 lines" in ws["B4"].value


class TestTheFlashNeverFailsBecauseOfTheBillingTab:
    def test_no_invoice_lines_writes_the_title_and_a_note(self, tmp_path):
        """Month-start quiet days and fresh deployments: the tab is there, says
        which date it is for, and says plainly that there is nothing on it."""
        quiet = [
            _inv_line(
                "SL00004",
                1013603,
                "A2",
                "CUSTOMER 2",
                10.0,
                90,
                10,
                100,
                day=date(2026, 7, 30),
            )
        ]
        out = build_flash.build(
            FLASH_DAY,
            None,
            str(tmp_path),
            data=(WB_HEADERS, _wb_rows()),
            billing_day=BILLING_DAY,
            inv_data=(INV_HEADERS, quiet),
        )
        wb = openpyxl.load_workbook(out)
        assert wb.sheetnames == ["Flash 03 Aug", "Billing 31 Jul", build_flash.ALL_TAB]
        ws = wb["Billing 31 Jul"]
        assert ws["B3"].value == "Billing Component Detail — Invoice date 31 Jul 2026"
        text = [str(c.value) for row in ws.iter_rows() for c in row if c.value]
        assert any("No invoice lines for 31 Jul 2026" in t for t in text)
        assert not any("GRAND TOTAL" in t for t in text)

    def test_an_unreadable_invoice_source_is_skipped_with_a_warning(
        self, tmp_path, capsys
    ):
        out = build_flash.build(
            FLASH_DAY,
            None,
            str(tmp_path),
            data=(WB_HEADERS, _wb_rows()),
            inv_file=str(tmp_path / "does-not-exist.xlsx"),
            billing_day=BILLING_DAY,
        )
        wb = openpyxl.load_workbook(out)
        assert wb.sheetnames == ["Flash 03 Aug", build_flash.ALL_TAB]
        err = capsys.readouterr().err
        assert "WARNING" in err and "Billing tab" in err and "31 Jul" in err

    def test_a_missing_column_in_the_invoice_export_is_skipped_too(
        self, tmp_path, capsys
    ):
        headers = [h for h in INV_HEADERS if h != "Doc Charge"]
        rows = [
            [v for h, v in zip(INV_HEADERS, r) if h != "Doc Charge"]
            for r in _inv_rows()
        ]
        out = build_flash.build(
            FLASH_DAY,
            None,
            str(tmp_path),
            data=(WB_HEADERS, _wb_rows()),
            billing_day=BILLING_DAY,
            inv_data=(headers, rows),
        )
        assert openpyxl.load_workbook(out).sheetnames == [
            "Flash 03 Aug",
            build_flash.ALL_TAB,
        ]
        assert "WARNING" in capsys.readouterr().err

    def test_the_billing_detail_itself_still_refuses_an_empty_day(self, tmp_path):
        """Unchanged behaviour: a not-yet-invoiced day gets a flash, never a
        Billing Detail — the quiet-day note belongs to the flash only."""
        with pytest.raises(SystemExit):
            build_billing_detail.build(
                BILLING_DAY,
                None,
                None,
                str(tmp_path),
                None,
                inv_data=(INV_HEADERS, []),
                credits_data=(CREDIT_HEADERS, []),
            )


class TestRunDailyPassesTheBillingArguments:
    def test_flash_branch_carries_inv_file_and_billing_date(
        self, tmp_path, monkeypatch
    ):
        import run_daily

        data_dir = tmp_path / "data"
        data_dir.mkdir()
        label = run_daily.fy_label(run_daily.fy_start())
        cred_start = run_daily.fy_start().year - 2
        for name in (
            f"WB Date - {label}..xlsx",
            f"INV Date - {label}..xlsx",
            f"Credits - March{cred_start % 100} - Feb{(run_daily.fy_start().year + 1) % 100}.xlsx",
            run_daily.PY_INV_NAME,
            run_daily.BUDGET_NAME,
        ):
            (data_dir / name).write_bytes(b"")

        calls = []
        monkeypatch.setattr(
            run_daily, "latest_dates", lambda wb, inv: (FLASH_DAY, BILLING_DAY)
        )
        monkeypatch.setattr(
            run_daily, "run", lambda step, args: calls.append((step, args)) or "x.xlsx"
        )
        run_daily.generate_reports(
            Namespace(
                data_dir=str(data_dir),
                report_dir=str(tmp_path / "out"),
                to=None,
                cc=None,
                skip_extract=True,
                only="flash",
                no_email=True,
                dry_run_email=False,
            )
        )
        ((step, args),) = calls
        assert step == "flash"
        assert args[0] == "build_flash.py"
        assert args[args.index("--date") + 1] == FLASH_DAY.isoformat()
        assert args[args.index("--inv-file") + 1] == str(
            data_dir / f"INV Date - {label}..xlsx"
        )
        assert args[args.index("--billing-date") + 1] == BILLING_DAY.isoformat()


class TestFlashCli:
    def test_billing_arguments_travel_together(self):
        """--inv-file without --billing-date (or the reverse) is a usage error
        at the command line; run_daily always passes both."""
        parser = build_flash.cli_parser()
        with pytest.raises(SystemExit):
            parser.parse_args(
                ["--date", "2026-08-03", "--wb-file", "x", "--inv-file", "y"]
            )
        with pytest.raises(SystemExit):
            parser.parse_args(
                [
                    "--date",
                    "2026-08-03",
                    "--wb-file",
                    "x",
                    "--billing-date",
                    "2026-07-31",
                ]
            )
        a = parser.parse_args(
            [
                "--date",
                "2026-08-03",
                "--wb-file",
                "x",
                "--inv-file",
                "y",
                "--billing-date",
                "2026-07-31",
            ]
        )
        assert a.inv_file == "y" and a.billing_date == "2026-07-31"
