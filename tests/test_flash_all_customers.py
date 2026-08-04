"""The flash's All Customers tab, and the page-one contract it must not break.

Larry asked on 4 Aug 2026 for the flash to show every client. The full list
went onto a second tab rather than into the front-page block, because
build_billing_detail._flash_comparison re-reads the frozen flash and finds its
sections by scanning the FIRST sheet for three literal headers. Lengthening or
renaming "TOP 10 CUSTOMERS" in place leaves the rep section open and pulls
every customer row into it — a hundred "reps" instead of six, with no error.

That parser also has to keep working against the archive of flashes frozen
before this change, so page one is fixed by contract, not by preference. These
tests encode both halves: the front page stays exactly as the parser expects,
and the new tab ties back to the KPI to the cent.
"""

import sys
from datetime import date
from pathlib import Path

import openpyxl
import pytest
import xlsxwriter

REPO = Path(__file__).resolve().parent.parent
BUILDERS = REPO / "revenue_reports"
sys.path.insert(0, str(BUILDERS))

import build_flash  # noqa: E402

# Headers _flash_comparison switches on. Copied deliberately rather than
# imported: they are string literals in the other builder, and the point of
# the test is to fail if the two drift apart.
PARSER_SECTIONS = ("BY BRANCH (origin)", "BY REP", "TOP 10 CUSTOMERS")

HEADERS = ["Waybill", "Waybill Date", "Customer", "Account", "Orig Hub",
           "Chrg Mass", "Subtotal", "Salesrep", "Rep"]

DAY = date(2026, 8, 3)
PRIOR = date(2026, 7, 27)          # same weekday, for the typical-Mon benchmark


def _row(wb, cust, acct, hub, kg, sub, rep, repname, day=DAY):
    return [wb, day, cust, acct, hub, kg, sub, rep, repname]


@pytest.fixture
def workbook(tmp_path):
    """Twelve accounts so the top ten is a genuine subset, plus one account
    carrying two spellings of its customer name — the case that grouping on
    Account rather than Customer exists to merge."""
    rows = []
    for i in range(12):
        rows.append(_row(f"SL{i:05d}", f"CUSTOMER {i}", f"A{i}", "JNB",
                         100 + i, 1000 - i * 10, "TF", "Tracy Flandorp"))
    # same account, second spelling, and a second rep and branch for coverage
    rows.append(_row("SL00100", "CUSTOMER 0 (PTY) LTD", "A0", "CPT",
                     50, 500, "CN", "Christine Naidoo"))
    # prior same-weekday history so `typical` has something to average
    rows.append(_row("SL00200", "CUSTOMER 1", "A1", "DUR", 10, 900, "TF",
                     "Tracy Flandorp", PRIOR))

    src = tmp_path / "wb.xlsx"
    w = xlsxwriter.Workbook(str(src))
    sh = w.add_worksheet()
    # A date-only number format, so the reader hands back date and not datetime
    # — build_flash compares Waybill Date to a date directly, as the real
    # Parcel Perfect export allows.
    dfmt = w.add_format({"num_format": "yyyy-mm-dd"})
    for c, h in enumerate(HEADERS):
        sh.write(0, c, h)
    for r, vals in enumerate(rows, start=1):
        for c, v in enumerate(vals):
            if isinstance(v, date):
                sh.write_datetime(r, c, v, dfmt)
            else:
                sh.write(r, c, v)
    w.close()

    out = build_flash.build(DAY, str(src), str(tmp_path))
    return openpyxl.load_workbook(out, data_only=True)


class TestPageOneStaysParseable:
    """The contract build_billing_detail._flash_comparison relies on."""

    def test_all_customers_is_not_the_first_sheet(self, workbook):
        # openpyxl's .active is what the other builder opens.
        assert workbook.active.title != build_flash.ALL_TAB
        assert workbook.sheetnames[0] == workbook.active.title

    def test_front_page_still_says_top_10_customers(self, workbook):
        col_b = [c.value for c in workbook.active["B"]]
        for header in PARSER_SECTIONS:
            assert header in col_b, f"{header!r} missing — the parser will mis-section"

    def test_rep_section_closes_before_the_customers(self, workbook):
        """Replays the parser and asserts it sees reps, not customers."""
        f, section, reps, branches = workbook.active, None, {}, {}
        for row in f.iter_rows(min_col=2, max_col=3):
            v = row[0].value
            if v in PARSER_SECTIONS:
                section = v
            elif section == "BY BRANCH (origin)" and v and row[1].value is not None and v != "Branch":
                branches[v] = row[1].value
            elif section == "BY REP" and v and row[1].value is not None and v != "Rep":
                reps[str(v).split(" — ")[0]] = row[1].value
        assert set(reps) == {"TF", "CN"}
        assert set(branches) == {"JHB", "Cape Town"}

    def test_by_rep_keeps_its_label_in_B_and_revenue_in_C(self, workbook):
        """Chg kg and R/kg were added to BY REP on 4 Aug at Larry's request.

        The flash comparison reads column B for the rep and column C for the
        revenue. Inserting the new columns to the LEFT of revenue would have
        fed it chargeable mass instead, with no error and a plausible-looking
        variance — so the column order is a contract, not a layout choice.
        """
        ws = workbook.active
        r = next(r for r in range(1, ws.max_row + 1)
                 if ws.cell(r, 2).value == "BY REP")
        assert [ws.cell(r + 1, c).value for c in range(2, 7)] == [
            "Rep", "Revenue", "Chg kg", "R/kg", "Waybills"]
        # first data row: revenue in C, and R/kg consistent with C / D
        rev, kg, rate = (ws.cell(r + 2, c).value for c in (3, 4, 5))
        assert rate == pytest.approx(rev / kg)

    def test_the_comparison_still_reads_revenue_not_mass(self, workbook):
        """Replays the other builder's parser and checks the figure it picks
        up for a rep is that rep's revenue."""
        ws = workbook.active
        f_rep, section = {}, None
        for row in ws.iter_rows(min_col=2, max_col=3):
            v = row[0].value
            if v in PARSER_SECTIONS:
                section = v
            elif section == "BY REP" and v and row[1].value is not None and v != "Rep":
                f_rep[str(v).split(" — ")[0]] = row[1].value
        # TF ships 12 waybills at 1000 down to 890, CN one at 500
        assert f_rep["CN"] == 500
        assert f_rep["TF"] == pytest.approx(sum(1000 - i * 10 for i in range(12)))

    def test_front_page_lists_ten_customers(self, workbook):
        ws = workbook.active
        start = next(r for r in range(1, ws.max_row + 1)
                     if ws.cell(r, 2).value == "TOP 10 CUSTOMERS")
        names = []
        for r in range(start + 2, ws.max_row + 1):
            v = ws.cell(r, 2).value
            if v is None or ws.cell(r, 3).value is None:
                break
            names.append(v)
        assert len(names) == 10


class TestAllCustomersTab:
    def test_lists_every_client_that_moved_freight(self, workbook):
        ws = workbook[build_flash.ALL_TAB]
        assert ws.cell(ws.max_row, 3).value == "12 clients"   # 13 rows, 12 accounts

    def test_merges_name_variants_on_one_account(self, workbook):
        ws = workbook[build_flash.ALL_TAB]
        accounts = [ws.cell(r, 2).value for r in range(8, ws.max_row)]
        assert len(accounts) == len(set(accounts)), "an account was split across rows"
        assert "A0" in accounts

    def test_total_reconciles_to_the_headline_revenue(self, workbook):
        """The old top-ten block never tied to the KPI; this one must."""
        kpi = workbook.active["B8"].value
        ws = workbook[build_flash.ALL_TAB]
        body = sum(ws.cell(r, 4).value for r in range(8, ws.max_row))
        assert body == pytest.approx(kpi, abs=0.005)
        assert ws.cell(ws.max_row, 4).value == pytest.approx(kpi, abs=0.005)

    def test_waybills_and_kg_reconcile_too(self, workbook):
        p1, ws = workbook.active, workbook[build_flash.ALL_TAB]
        assert ws.cell(ws.max_row, 7).value == p1["D8"].value          # waybills
        assert ws.cell(ws.max_row, 5).value == pytest.approx(p1["F8"].value)

    def test_cumulative_share_reaches_one(self, workbook):
        ws = workbook[build_flash.ALL_TAB]
        assert ws.cell(ws.max_row - 1, 9).value == pytest.approx(1.0)

    def test_rows_are_ordered_by_revenue(self, workbook):
        ws = workbook[build_flash.ALL_TAB]
        rev = [ws.cell(r, 4).value for r in range(8, ws.max_row)]
        assert rev == sorted(rev, reverse=True)

    def test_prints_with_repeating_headings(self, workbook):
        """~100 rows a day — it gets printed, not just scrolled."""
        ws = workbook[build_flash.ALL_TAB]
        assert ws.print_title_rows == "$7:$7"
        assert ws.page_setup.orientation == "landscape"


class TestMatchesTheHouseFormatting:
    """Read off Larry's "Flash Revenue - 27 Jul 2026" reference and the billing
    detail's "By Customer" tab — the other long per-customer list in the suite.

    Formatting drift is the failure mode this project has actually had: five
    passes over four days in early August restored tab colours, freezes, row
    heights, fills, fonts and borders, none of which diff_workbooks.py could
    see because it compares values. So the conventions get asserted.
    """

    def test_tab_is_coloured_like_the_rest_of_the_workbook(self, workbook):
        ws = workbook[build_flash.ALL_TAB]
        assert ws.sheet_properties.tabColor.rgb.endswith("FF6900")   # ORANGE

    def test_headings_are_navy2_as_the_flash_reference_has_them(self, workbook):
        """NAVY2, not the NAVY the billing detail uses — this sheet lives in
        the flash workbook, whose every heading band is NAVY2."""
        ws = workbook[build_flash.ALL_TAB]
        for c in range(2, 10):
            cell = ws.cell(7, c)
            assert cell.fill.start_color.rgb.endswith("0A0050"), cell.value
            assert cell.font.bold and cell.font.size == 9

    def test_heading_row_carries_the_house_height(self, workbook):
        """H_BILLING. Left unset it falls back to Excel's 15 and renders
        noticeably tighter than Larry's — the exact drift style.py warns about."""
        ws = workbook[build_flash.ALL_TAB]
        assert ws.row_dimensions[7].height == pytest.approx(21.9)

    def test_headings_stay_visible_when_the_detail_scrolls(self, workbook):
        assert workbook[build_flash.ALL_TAB].freeze_panes == "B8"

    def test_figures_are_blue_labels_black_and_rows_stripe(self, workbook):
        ws = workbook[build_flash.ALL_TAB]
        assert ws.cell(8, 4).font.color.rgb.endswith("0000FF")     # revenue
        assert ws.cell(8, 3).font.color.rgb.endswith("000000")     # customer name
        assert ws.cell(8, 3).fill.start_color.rgb.endswith("FFFFFF")
        assert ws.cell(9, 3).fill.start_color.rgb.endswith("F0F0F8")   # ALT
        assert ws.cell(8, 3).font.size == 9        # the long-list size, not 10

    def test_every_populated_cell_is_boxed(self, workbook):
        """The reference draws a thin D9D9D9 grid round every populated cell."""
        ws = workbook[build_flash.ALL_TAB]
        for r in (7, 8, ws.max_row):
            for c in range(2, 10):
                assert ws.cell(r, c).border.left.style == "thin", f"r{r}c{c}"

    def test_number_formats_come_from_the_shared_constants(self, workbook):
        from style import DEC2, NUM, PCT1
        ws = workbook[build_flash.ALL_TAB]
        assert ws.cell(8, 4).number_format == NUM      # revenue
        assert ws.cell(8, 6).number_format == DEC2     # R/kg
        assert ws.cell(8, 8).number_format == PCT1     # % of day

    def test_total_row_is_the_house_orange_band(self, workbook):
        ws = workbook[build_flash.ALL_TAB]
        cell = ws.cell(ws.max_row, 2)
        assert cell.fill.start_color.rgb.endswith("FF6900") and cell.font.bold


class TestTheRuntimeGuard:
    """The contract above is enforced at build time, not just in CI.

    The tests in this file protect the flash's layout. They cannot protect a
    frozen flash already sitting in the reports folder, or a change made to the
    flash by someone who never runs pytest. _check_flash_parse re-derives the
    invariant from whatever it actually read: every branch and every rep with
    revenue appears in the flash, so each block must sum to the headline.

    It warns rather than raises. A wrong comparison tab is worth flagging
    loudly; it is not worth withholding four correct reports over.
    """

    def _warnings(self, capsys, f_rev, f_branch, f_rep):
        import build_billing_detail
        build_billing_detail._check_flash_parse("flash.xlsx", f_rev, f_branch, f_rep)
        return capsys.readouterr().err

    def test_silent_when_the_blocks_reconcile(self, capsys):
        err = self._warnings(capsys, 1000.0,
                             {"JHB": 600.0, "Durban": 400.0},
                             {"TF": 700.0, "CN": 300.0})
        assert err == ""

    def test_flags_mass_read_as_revenue(self, capsys):
        """Chg kg inserted left of Revenue: every figure is plausible, and the
        block no longer sums to the day."""
        err = self._warnings(capsys, 1000.0,
                             {"JHB": 600.0, "Durban": 400.0},
                             {"TF": 140.0, "CN": 60.0})
        assert "BY REP" in err and "no longer revenue" in err
        assert "BY BRANCH" not in err, "the branch block was fine and must not warn"

    def test_flags_a_renamed_section_header(self, capsys):
        err = self._warnings(capsys, 1000.0, {"JHB": 600.0, "Durban": 400.0}, {})
        assert "read no rows" in err and "renamed" in err

    def test_tolerates_rounding_but_not_a_real_gap(self, capsys):
        assert self._warnings(capsys, 1000.0, {"JHB": 1000.4}, {"TF": 999.7}) == ""
        assert "no longer revenue" in self._warnings(
            capsys, 1000.0, {"JHB": 1000.0}, {"TF": 990.0})

    def test_says_nothing_when_there_is_no_headline_to_compare(self, capsys):
        """A flash with an empty KPI band is a different fault; do not add a
        second, misleading warning on top of it."""
        assert self._warnings(capsys, 0, {"JHB": 5.0}, {"TF": 5.0}) == ""
