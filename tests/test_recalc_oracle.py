"""Check our cached formula values against an independent spreadsheet engine.

Every formula the builders write carries a value we computed in Python. If that
arithmetic is wrong the file is worse than before — a wrong number looks right,
where a zero at least looked broken. So the values are checked against
LibreOffice's own evaluation of the same formulas.

Opt-in, because it needs LibreOffice and the real export files:

    pytest tests/ -m oracle --workbook "/path/to/Revenue Dashboard.xlsx"

LibreOffice preserves cached values on load by default, so a profile that
forces recalculation is created first — without it the oracle would simply
read our own numbers back and always agree.
"""

import os
import shutil
import subprocess
import tempfile
from pathlib import Path

import pytest

pytestmark = pytest.mark.oracle

RECALC_ALWAYS = (
    '<item oor:path="/org.openoffice.Office.Calc/Formula/Load">'
    '<prop oor:name="OOXMLRecalcMode" oor:op="fuse"><value>0</value></prop></item>'
    '<item oor:path="/org.openoffice.Office.Calc/Formula/Load">'
    '<prop oor:name="ODFRecalcMode" oor:op="fuse"><value>0</value></prop></item>'
)


def _profile(tmp: Path) -> dict:
    """A LibreOffice user profile set to recalculate on load."""
    env = dict(os.environ, HOME=str(tmp))
    subprocess.run(["soffice", "--headless", "--terminate_after_init"],
                   env=env, capture_output=True, timeout=180)
    for reg in tmp.rglob("registrymodifications.xcu"):
        s = reg.read_text()
        if "OOXMLRecalcMode" not in s:
            reg.write_text(s.replace("</oor:items>", RECALC_ALWAYS + "</oor:items>"))
    return env


def recalculate(path: Path, tmp: Path):
    """Return LibreOffice's evaluation of every cell in the workbook."""
    import openpyxl
    out = tmp / "recalc"
    out.mkdir(exist_ok=True)
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx",
                    "--outdir", str(out), str(path)],
                   env=_profile(tmp), check=True, capture_output=True, timeout=900)
    return openpyxl.load_workbook(out / path.name, data_only=True)


def test_cached_values_match_libreoffice(request):
    workbook = request.config.getoption("--workbook")
    if not workbook:
        pytest.skip("pass --workbook to run the oracle")
    if not shutil.which("soffice"):
        pytest.skip("LibreOffice not installed")
    import openpyxl

    path = Path(workbook)
    ours = openpyxl.load_workbook(path, data_only=True)
    formulas = openpyxl.load_workbook(path)

    with tempfile.TemporaryDirectory() as td:
        theirs = recalculate(path, Path(td))

        checked, bad = 0, []
        for name in formulas.sheetnames:
            a, b, f = ours[name], theirs[name], formulas[name]
            for row in f.iter_rows():
                for cell in row:
                    if not (isinstance(cell.value, str) and cell.value.startswith("=")):
                        continue
                    checked += 1
                    x, y = a[cell.coordinate].value, b[cell.coordinate].value
                    if isinstance(x, (int, float)) and isinstance(y, (int, float)):
                        if abs(x - y) > 0.51:
                            bad.append(f"{name}!{cell.coordinate} ours={x!r} calc={y!r}")
                    elif (x is None or x == "") and (y is None or y == ""):
                        continue          # both blank — the IF(...,"") guards
                    elif x != y:
                        bad.append(f"{name}!{cell.coordinate} ours={x!r} calc={y!r}")

    assert checked, "no formulas found — is this the right workbook?"
    assert not bad, f"{len(bad)} of {checked} disagree:\n" + "\n".join(bad[:40])
