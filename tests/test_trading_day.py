"""Picking the day a report covers.

"Latest date in the file" is not "last day that traded", and the difference
sent exco a flash reporting R0. On Monday 3 Aug 2026 the waybill export held
ten waybills dated Sunday the 2nd worth nothing at all, and the invoice export
held two lines dated Friday the 31st because that day's invoice run had not
happened yet. Taking the max picked both.

Second occurrence, 20 Oct 2026: the value floor alone was not enough. October's
invoicing landed in batches (9 Oct R5.48m, 14 Oct R3.54m, 16 Oct R1.95m) with
five-to-seven-line days between them, and because the median was taken over
every day with ANY rows, those thin days dragged the floor down to R8,075. A
Monday with three invoice lines (R16,240) then cleared it, and the 16 Oct PM
pack had already shipped a five-line Billing Detail for 15 Oct. So a day now
also needs MIN_WAYBILLS rows to count, and the median is taken over such days
only. The pairs contract grew a third element for that: (date, value) is one
row; (date, value, rows) is an already-aggregated day.
"""

import sys
from datetime import date, timedelta
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "revenue_reports"))

import day_guards
from run_daily import last_trading_day

MON = date(2026, 8, 3)
FULL = 700  # rows on an ordinary day; every real day in these tests is this size


def week(start: date, values, counts=None):
    """One aggregated (date, value, rows) item per day. Rows default to a
    full day's worth, so the value floor is what each case exercises."""
    counts = counts or [FULL] * len(values)
    return [
        (start + timedelta(days=i), v, n)
        for i, (v, n) in enumerate(zip(values, counts))
    ]


class TestRealCase:
    """The actual figures from 3 August 2026."""

    def test_waybills_roll_back_over_the_weekend(self):
        pairs = week(
            date(2026, 7, 27),
            [
                632846,  # Mon 27
                728287,  # Tue 28
                803604,  # Wed 29
                637612,  # Thu 30
                717045,  # Fri 31
                78938,  # Sat 1 Aug — 11% of a normal day
                0,  # Sun 2 Aug — ten waybills, no value
            ],
            [FULL, FULL, FULL, FULL, FULL, 80, 10],
        )
        assert last_trading_day(pairs, MON) == date(2026, 7, 31)

    def test_invoices_skip_the_day_whose_run_has_not_happened(self):
        pairs = week(
            date(2026, 7, 27),
            [
                819559,  # Mon 27
                684551,  # Tue 28
                830847,  # Wed 29
                638341,  # Thu 30
                840,  # Fri 31 — two lines, invoicing not yet run
            ],
            [FULL, FULL, FULL, FULL, 2],
        )
        assert last_trading_day(pairs, MON) == date(2026, 7, 30)


class TestThinDaysAreNotTradingDays:
    """The 20 Oct 2026 shapes, figures straight from the 09:55 INV extract."""

    OCT = [
        (date(2026, 10, 8), 951, 1),
        (date(2026, 10, 9), 5_480_835, 5702),
        (date(2026, 10, 12), 8_584, 7),
        (date(2026, 10, 13), 4_872, 7),
        (date(2026, 10, 14), 3_539_210, 3461),
        (date(2026, 10, 15), 14_137, 5),
        (date(2026, 10, 16), 1_945_212, 1901),
        (date(2026, 10, 19), 16_240, 3),
    ]
    # the batch days that form the real floor: late September invoicing
    SEP = [(date(2026, 9, d), 2_000_000, 1500) for d in (22, 23, 24, 25, 28, 29, 30)]

    def test_three_lines_on_a_monday_are_not_a_billing_day(self):
        """The 09:55 dry run picked 19 Oct. It must pick 16 Oct."""
        assert last_trading_day(self.SEP + self.OCT, date(2026, 10, 20)) == date(
            2026, 10, 16
        )

    def test_five_lines_are_not_a_billing_day_either(self):
        """The 16 Oct PM pack shipped a Billing Detail for 15 Oct with 5 lines."""
        pairs = self.SEP + [p for p in self.OCT if p[0] < date(2026, 10, 16)]
        assert last_trading_day(pairs, date(2026, 10, 16)) == date(2026, 10, 14)

    def test_thin_days_do_not_drag_the_floor_down(self):
        """The failure mechanism itself: the median must be taken over days
        with MIN_WAYBILLS rows, so the floor is in the millions, not R8k."""
        _, _, floor = day_guards._totals_and_floor(
            self.SEP + self.OCT, date(2026, 10, 20)
        )
        assert floor > 500_000, (
            f"floor {floor:,.0f} — thin days are still in the median"
        )

    def test_the_same_shape_as_per_row_pairs(self):
        """run_daily hands over one pair per export row. Expanding each day
        into that many rows must give the same answer as the aggregated form."""
        expanded = [(d, v / n) for d, v, n in self.SEP + self.OCT for _ in range(n)]
        assert last_trading_day(expanded, date(2026, 10, 20)) == date(2026, 10, 16)

    def test_is_completed_trading_day_agrees(self):
        pairs, today = self.SEP + self.OCT, date(2026, 10, 20)
        assert (
            day_guards.is_completed_trading_day(pairs, date(2026, 10, 16), today)
            is True
        )
        assert (
            day_guards.is_completed_trading_day(pairs, date(2026, 10, 19), today)
            is False
        )
        assert (
            day_guards.is_completed_trading_day(pairs, date(2026, 10, 15), today)
            is False
        )

    def test_a_quiet_but_real_day_with_enough_lines_still_counts(self):
        """Month-start: 25 lines at 60% of the median is a billing day, as it
        was before — the row floor adds a condition, it does not replace the
        value test."""
        pairs = self.SEP + [(date(2026, 10, 1), 1_200_000, 25)]
        assert last_trading_day(pairs, date(2026, 10, 2)) == date(2026, 10, 1)

    def test_enough_lines_but_almost_no_value_is_still_skipped(self):
        """And the value floor still bites on its own: 25 lines worth R20k is
        a mis-keyed handful, not a day's invoicing."""
        pairs = self.SEP + [(date(2026, 10, 1), 20_000, 25)]
        assert last_trading_day(pairs, date(2026, 10, 2)) == date(2026, 9, 30)

    def test_one_row_short_of_the_minimum_is_skipped(self):
        pairs = self.SEP + [(date(2026, 10, 1), 1_900_000, day_guards.MIN_WAYBILLS - 1)]
        assert last_trading_day(pairs, date(2026, 10, 2)) == date(2026, 9, 30)
        pairs = self.SEP + [(date(2026, 10, 1), 1_900_000, day_guards.MIN_WAYBILLS)]
        assert last_trading_day(pairs, date(2026, 10, 2)) == date(2026, 10, 1)


class TestOrdinaryDays:
    def test_midweek_picks_yesterday(self):
        pairs = week(date(2026, 7, 27), [600000, 600000, 600000, 620000])
        assert last_trading_day(pairs, date(2026, 7, 31)) == date(2026, 7, 30)

    def test_today_is_never_chosen(self):
        pairs = week(date(2026, 7, 27), [600000, 600000, 600000, 600000, 900000])
        assert last_trading_day(pairs, date(2026, 7, 31)) == date(2026, 7, 30)

    def test_a_quiet_but_real_day_still_counts(self):
        # Half the median is the floor, so a day at 60% is a trading day.
        pairs = week(date(2026, 7, 27), [600000, 600000, 600000, 600000, 360000])
        assert last_trading_day(pairs, MON) == date(2026, 7, 31)

    def test_a_public_holiday_is_skipped(self):
        pairs = week(date(2026, 7, 27), [600000, 600000, 600000, 600000, 12000])
        assert last_trading_day(pairs, MON) == date(2026, 7, 30)


class TestDegenerate:
    def test_empty_input_returns_today(self):
        assert last_trading_day([], MON) == MON

    def test_single_day_of_history_is_used_rather_than_rejected(self):
        assert last_trading_day([(date(2026, 7, 31), 5000)], MON) == date(2026, 7, 31)
        assert last_trading_day([(date(2026, 7, 31), 5000, FULL)], MON) == date(
            2026, 7, 31
        )

    def test_no_day_big_enough_to_judge_falls_back_to_the_newest(self):
        """A sparse file — every day under MIN_WAYBILLS — has no floor to
        compare against, so the newest day is used rather than nothing."""
        pairs = [(date(2026, 7, 30), 900, 3), (date(2026, 7, 31), 500, 2)]
        assert last_trading_day(pairs, MON) == date(2026, 7, 31)
        assert (
            day_guards.is_completed_trading_day(pairs, date(2026, 7, 31), MON) is True
        )

    def test_all_days_below_the_floor_falls_back_to_the_newest(self):
        """Never return nothing — a wrong day beats no report at all."""
        pairs = [(date(2026, 7, 30), 0), (date(2026, 7, 31), 0)]
        assert last_trading_day(pairs, MON) == date(2026, 7, 31)
        pairs = [(date(2026, 7, 30), 0, FULL), (date(2026, 7, 31), 0, FULL)]
        assert last_trading_day(pairs, MON) == date(2026, 7, 31)

    def test_none_values_are_tolerated(self):
        pairs = [(date(2026, 7, 30), None), (date(2026, 7, 31), 600000)]
        assert last_trading_day(pairs, MON) == date(2026, 7, 31)
        pairs = [(date(2026, 7, 30), None, FULL), (date(2026, 7, 31), 600000, FULL)]
        assert last_trading_day(pairs, MON) == date(2026, 7, 31)

    def test_non_dates_are_ignored(self):
        pairs = [("not a date", 999999), (date(2026, 7, 31), 600000)]
        assert last_trading_day(pairs, MON) == date(2026, 7, 31)

    def test_a_bare_pair_counts_as_one_row(self):
        """The contract run_daily relies on: one (date, value) per export row."""
        totals, counts, _ = day_guards._totals_and_floor(
            [(MON - timedelta(days=1), 10)] * 25 + [(MON - timedelta(days=2), 10, 300)],
            MON,
        )
        assert counts[MON - timedelta(days=1)] == 25
        assert counts[MON - timedelta(days=2)] == 300
        assert totals[MON - timedelta(days=1)] == 250


class TestUnbilledBillingFrontier:
    """End-to-end: does the report itself exclude normal billing lag?

    build_unbilled derives a billing frontier — the first day whose invoicing
    has not run — so freight still waiting for a normal invoice run is not
    reported as unbilled. Two things have broken it, and this class drives the
    whole builder rather than the frontier function, so it fails if either
    returns and if the wiring between them ever comes loose.

    Mis-keyed years put waybills in 2522 and 9473, two marked Invoiced, which
    pushed a plain max() into the year 9473 and excluded nothing: 1,050
    waybills and R3.07m against 35-85 and R10-265k in Larry's own reports.

    Then one real, current, correctly-dated early invoice moved that same max
    and took the report from R41k to R731k with nothing about the day changed.

    Volumes here are realistic on purpose. An earlier version of this class
    used four rows, which sits under MIN_WAYBILLS, so every case quietly ran
    down the fallback branch and the main rule was never executed at all.
    """

    HEADERS = [
        "Waybill",
        "Waybill Date",
        "Account",
        "Customer",
        "Service",
        "Status",
        "Subtotal",
    ]

    JUL30, JUL31 = date(2026, 7, 30), date(2026, 7, 31)
    AUG3 = date(2026, 8, 3)

    def _build(self, tmp_path, rows):
        import build_unbilled
        import openpyxl
        import xlsxwriter

        tmp_path.mkdir(parents=True, exist_ok=True)
        src = tmp_path / "wb.xlsx"
        w = xlsxwriter.Workbook(str(src))
        sh = w.add_worksheet()
        dfmt = w.add_format({"num_format": "yyyy-mm-dd"})
        for c, h in enumerate(self.HEADERS):
            sh.write(0, c, h)
        for r, vals in enumerate(rows, start=1):
            for c, v in enumerate(vals):
                if isinstance(v, date):
                    sh.write_datetime(r, c, v, dfmt)
                else:
                    sh.write(r, c, v)
        w.close()
        out = build_unbilled.build(str(src), str(tmp_path))
        return openpyxl.load_workbook(out, data_only=True)

    def _wb(self, n, day, status, value, tag):
        return [
            [f"{tag}{i:04d}", day, f"A{tag}", f"C{tag}", "RDF", status, value]
            for i in range(n)
        ]

    def _rows(self, *, typo=False, aug3_invoiced=0, aug3_total=100):
        """July invoiced and done; August shipped but not yet invoiced.

        R500 of genuinely unbilled freight sits inside July — that is the only
        figure the report should ever report here. August is billing lag.
        """
        rows = []
        rows += self._wb(50, self.JUL30, "Invoiced", 100, "P")
        rows += self._wb(50, self.JUL31, "Invoiced", 100, "Q")
        rows += self._wb(1, date(2026, 7, 29), "Ready for Approval", 500, "U")
        rows += self._wb(aug3_invoiced, self.AUG3, "Invoiced", 90, "V")
        rows += self._wb(
            aug3_total - aug3_invoiced, self.AUG3, "Ready for Approval", 90, "W"
        )
        if typo:
            # straight from the export: an Invoiced waybill keyed to year 9473
            rows += self._wb(1, date(9473, 7, 5), "Invoiced", 100, "Z")
        return rows

    def _headline(self, wb):
        ws = wb["Overview"]
        for r in range(1, 12):
            if str(ws.cell(r, 1).value or "").startswith("Unbilled value"):
                return ws.cell(r, 3).value

    def test_august_is_billing_lag_not_unbilled_freight(self, tmp_path):
        """The baseline: only the July waybill counts."""
        assert self._headline(self._build(tmp_path, self._rows())) == 500

    def test_a_year_9473_invoice_does_not_disable_the_frontier(self, tmp_path):
        clean = self._headline(self._build(tmp_path / "a", self._rows()))
        typo = self._headline(self._build(tmp_path / "b", self._rows(typo=True)))
        assert clean == typo == 500, (
            "a waybill dated 9473 must not move the frontier — that is the "
            "defect that reported R3.07m on 3 Aug 2026"
        )

    def test_one_early_august_invoice_does_not_move_the_frontier(self, tmp_path):
        """The 10:40 case: real, current, correctly dated, and it moved a max."""
        wb = self._build(tmp_path, self._rows(aug3_invoiced=1))
        assert self._headline(wb) == 500, (
            "one invoice is not evidence that 3 August has been billed; "
            "counting it drags all 99 remaining waybills in as unbilled"
        )

    def test_the_frontier_moves_once_august_is_really_invoiced(self, tmp_path):
        """And it must still move when the day genuinely is done — 60 of 100
        invoiced leaves 40 at R90 truly unbilled, plus the July 500."""
        wb = self._build(tmp_path, self._rows(aug3_invoiced=60))
        assert self._headline(wb) == 500 + 40 * 90

    def test_the_lagging_waybills_are_the_ones_excluded(self, tmp_path):
        """Guards the assertions above against passing for the wrong reason."""
        wb = self._build(tmp_path, self._rows(typo=True, aug3_invoiced=1))
        detail = wb["Unbilled Detail"]
        # Skip the trailing TOTAL row, which carries no waybill number.
        listed = {
            v for r in range(2, detail.max_row + 1) if (v := detail.cell(r, 2).value)
        }
        assert "U0000" in listed, "the genuinely unbilled July waybill is missing"
        assert not any(w.startswith("W") for w in listed), (
            "August waybills awaiting their invoice run must not be listed"
        )
        assert not any(w.startswith("Z") for w in listed), (
            "the year-9473 waybill must not be listed"
        )


class TestExtractionIsBoundedAtBothEnds:
    """The manual export names a date RANGE; this query only had a floor.

    Parcel Perfect holds waybills with mis-keyed dates — 340 in the FY27 pull
    on 4 Aug 2026, years 2520 to 9473, the two Invoiced ones on records
    captured in October 2008. With no ceiling they all read as FY27. Larry's
    own export never showed one: a waybill dated 9473 is outside March26-Feb27.
    """

    def test_both_bounds_are_in_the_query(self):
        from extract_revenue import extraction_sql

        sql, params = extraction_sql("wb", date(2026, 3, 1))
        assert "wba.WAYDATE >= ?" in sql
        assert "wba.WAYDATE < ?" in sql
        assert params == [date(2026, 3, 1), date(2027, 3, 1)]

    def test_the_ceiling_closes_the_financial_year_that_start_opens(self):
        from extract_revenue import fy_end

        assert fy_end(date(2026, 3, 1)) == date(2027, 3, 1)
        assert fy_end(date(2025, 3, 1)) == date(2026, 3, 1)

    def test_invoice_basis_is_bounded_the_same_way(self):
        from extract_revenue import extraction_sql

        sql, params = extraction_sql("inv", date(2026, 3, 1))
        assert "wba.INVDATE >= ?" in sql
        assert "wba.INVDATE < ?" in sql
        assert params == [date(2026, 3, 1), date(2027, 3, 1)]


class TestBillingFrontierIsNotAPlainMax:
    """One invoice must not decide that a day is billed.

    On 4 Aug 2026 invoicing had not run for August at all — 0 of 807 waybills
    dated the 3rd were invoiced in the 07:13 export. By 10:40 a single early
    invoice carrying a 3 Aug waybill date had moved max() from 31 Jul to 3 Aug,
    which reclassified all of Monday's freight as unbilled and took the report
    from R41k to R731k. Nothing about the day had changed.

    Fixing the extract's date range does not help here: this row is real,
    current and correctly dated. Only judging a day by the SHARE invoiced does.
    """

    IX = {"Waybill Date": 0, "Status": 1}

    @staticmethod
    def _norm(v):
        return v

    def _rows(self, aug3_invoiced: int):
        """July fully invoiced, August not — plus however many early invoices
        happen to carry a 3 August waybill date."""
        rows = []
        for day in (30, 31):
            rows += [[date(2026, 7, day), "Invoiced"] for _ in range(700)]
            rows += [[date(2026, 7, day), "Ready for Approval"] for _ in range(10)]
        rows += [[date(2026, 8, 3), "Invoiced"] for _ in range(aug3_invoiced)]
        rows += [
            [date(2026, 8, 3), "Ready for Approval"] for _ in range(807 - aug3_invoiced)
        ]
        return rows

    def _frontier(self, rows):
        import build_unbilled

        return build_unbilled.billing_frontier(rows, self.IX, self._norm)

    def test_one_early_invoice_does_not_move_the_frontier(self):
        assert self._frontier(self._rows(aug3_invoiced=1)) == date(2026, 8, 1)

    def test_nor_does_a_handful(self):
        assert self._frontier(self._rows(aug3_invoiced=40)) == date(2026, 8, 1)

    def test_the_frontier_moves_once_the_day_is_actually_invoiced(self):
        assert self._frontier(self._rows(aug3_invoiced=500)) == date(2026, 8, 4)

    def test_a_future_dated_invoice_is_still_rejected(self):
        """The year-9473 case, kept from the earlier fix."""
        rows = self._rows(aug3_invoiced=0)
        rows += [[date(9473, 7, 5), "Invoiced"]] * 50
        assert self._frontier(rows) == date(2026, 8, 1)

    def test_a_quiet_day_is_too_small_to_judge(self):
        """A Saturday carrying three waybills, all invoiced, is not evidence
        that invoicing has reached Saturday."""
        rows = self._rows(aug3_invoiced=0)
        rows += [[date(2026, 8, 8), "Invoiced"]] * 3
        assert self._frontier(rows) == date(2026, 8, 1)


class TestBillingFrontierFallback:
    """The branch taken when no day is big enough to judge.

    It exists so a first run, or a sparse file, still produces a report rather
    than failing. It is the OLD rule, so it must keep the future-date guard —
    otherwise a year-9473 row walks straight back in through the back door.

    Covered explicitly because the end-to-end tests deliberately use realistic
    volumes now, which means nothing else reaches this branch.
    """

    IX = {"Waybill Date": 0, "Status": 1}

    @staticmethod
    def _norm(v):
        return v

    def _frontier(self, rows):
        import build_unbilled

        return build_unbilled.billing_frontier(rows, self.IX, self._norm)

    def test_a_sparse_file_still_yields_a_frontier(self):
        rows = [
            [date(2026, 7, 31), "Invoiced"],
            [date(2026, 8, 3), "Ready for Approval"],
        ]
        assert self._frontier(rows) == date(2026, 8, 1)

    def test_the_fallback_still_rejects_future_dates(self):
        rows = [[date(2026, 7, 31), "Invoiced"], [date(9473, 7, 5), "Invoiced"]]
        assert self._frontier(rows) == date(2026, 8, 1), (
            "the fallback must not reintroduce the year-9473 defect"
        )

    def test_no_invoiced_waybills_at_all_is_an_error_not_a_guess(self):
        import pytest

        with pytest.raises(SystemExit):
            self._frontier([[date(2026, 8, 3), "Ready for Approval"]])


class TestSaCalendar:
    """Trading days = weekdays minus SA public holidays (Larry, 13 Aug 2026).

    The 12 Aug 2026 dashboard projected August off "7 of 21 trading days",
    counting Mon 10 Aug — Women's Day observed — as a trading day the flash
    itself had skipped as non-trading. Larry: "we should be on day 6".
    """

    def test_august_2026_counts_20_days_not_21(self):
        from sa_calendar import trading_days_between

        assert trading_days_between(date(2026, 8, 1), date(2026, 8, 31)) == 20

    def test_elapsed_through_11_aug_is_larrys_day_6(self):
        from sa_calendar import trading_days_between

        assert trading_days_between(date(2026, 8, 1), date(2026, 8, 11)) == 6

    def test_womens_day_sunday_observance_lands_on_monday_10th(self):
        from sa_calendar import public_holidays

        assert date(2026, 8, 9) in public_holidays(2026)
        assert date(2026, 8, 10) in public_holidays(2026)

    def test_fy27_month_counts(self):
        # Mar 22 (Human Rights Day is a Saturday), Apr 19 (Good Friday,
        # Family Day, Freedom Day), May 20 (Workers' Day), Jun 21 (Youth
        # Day), Jul 23 (no weekday holidays).
        from build_dashboard import month_trading_days

        assert [month_trading_days(2026, m) for m in (3, 4, 5, 6, 7, 8)] == [
            22,
            19,
            20,
            21,
            23,
            20,
        ]

    def test_easter_2026(self):
        from sa_calendar import public_holidays

        hol = public_holidays(2026)
        assert date(2026, 4, 3) in hol  # Good Friday 2026
        assert date(2026, 4, 6) in hol  # Family Day 2026

    def test_prior_year_lfl_window_is_holiday_aware(self):
        # 6th trading day of Aug 2025 (1 Aug 2025 is a Friday; Women's Day
        # falls on a Saturday, so no weekday holiday that month).
        from sa_calendar import nth_trading_day

        assert nth_trading_day(2025, 8, 6) == date(2025, 8, 8)
        # April 2025: Freedom Day (Sun 27th) observed Mon 28th, plus Good
        # Friday 18th and Family Day 21st — 19 trading days.
        from sa_calendar import trading_days_between

        assert trading_days_between(date(2025, 4, 1), date(2025, 4, 30)) == 19

    def test_extra_workday_and_extra_holiday_overrides(self):
        import sa_calendar
        from sa_calendar import is_trading_day

        sat = date(2026, 8, 15)
        shutdown = date(2026, 8, 21)
        try:
            sa_calendar.EXTRA_WORKDAYS.add(sat)
            sa_calendar.EXTRA_HOLIDAYS.add(shutdown)
            sa_calendar.public_holidays.cache_clear()
            assert is_trading_day(sat)
            assert not is_trading_day(shutdown)
        finally:
            sa_calendar.EXTRA_WORKDAYS.discard(sat)
            sa_calendar.EXTRA_HOLIDAYS.discard(shutdown)
            sa_calendar.public_holidays.cache_clear()
