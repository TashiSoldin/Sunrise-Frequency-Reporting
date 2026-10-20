"""Format pin for the shared Billing tab writer.

Fable review finding 3 (20 Oct 2026, MEDIUM): the only automated format check
on build_billing_detail.write_billing_tab compared the flash's copy of the tab
with a Billing Detail written by the SAME function — a tautology. Two mutations
inside the writer (R/kg number format DEC2 → NUM; Customer column width 26 →
10) left every test green. The refactor's zero-diff proof against the 16 Oct
Billing Detail was a one-off manual check.

These tests pin the layout to absolute values read off Larry's 24 Jul 2026
Billing Detail reference (via the builder that reproduces it to the cell), so a
change to the writer fails here — once, for both reports, since both call the
same function.
"""

import sys
from datetime import date
from pathlib import Path

import openpyxl
import pytest
import xlsxwriter

REPO = Path(__file__).resolve().parent.parent
BUILDERS = REPO / "revenue_reports"
sys.path.insert(0, str(BUILDERS))

import build_billing_detail
from style import ALT, DEC2, NAVY, NUM, NUM1, NUM2, ORANGE, YELLOW, Styles

DAY = date(2026, 7, 31)

INV_HEADERS = [
    "Waybill", "Waybill Date", "Invoice", "Account", "Customer", "Shipper",
    "Consignee", "Orig Hub", "Dest Hub", "Dest Place", "Reference", "Service",
    "Pieces", "Actual Mass", "Chrg Mass", "Basic Charge", "Outlying Charge",
    "Doc Charge", "Fuel", "Subtotal", "Invoice Date", "Salesrep", "Rep",
    "First Ref", "Consolidated",
]  # fmt: skip


def _line(wb, acc, cust, kg, sub):
    return [
        wb, date(2026, 7, 29), 1013601.0, acc, cust, f"{cust} shipper", "Consignee",
        "JNB", "CPT", "Cape Town", f"REF-{wb}", "ONX", 2, kg - 0.5, kg,
        sub * 0.8, 0, 15, sub * 0.12, sub, DAY, "TF", "Tracy Flandorp", "", "No",
    ]  # fmt: skip


ROWS = [
    _line("SL00001", "A1", "CUSTOMER 1", 120.0, 1000),
    _line("SL00002", "A1", "CUSTOMER 1", 60.0, 500),
    _line("SL00003", "A0", "CUSTOMER 0", 20.0, 200),
]

# The 23 widths the writer sets, A..W, as they are in the reference.
WIDTHS = [2, 11, 9.5, 10, 8, 26, 22, 22, 9.5, 9.5, 18, 12, 9.5, 9.5, 9.5, 9.5, 9.5, 9.5, 9.5, 9.5, 9.5, 11, 8]  # fmt: skip

# Column letter -> number format on a data row. Text columns carry none.
NUMBER_FORMATS = {
    "N": NUM,  # Pcs
    "O": NUM1,  # Actual Mass
    "P": NUM1,  # Chrge Mass
    "Q": NUM2,  # Basic
    "R": NUM2,  # Outly
    "S": NUM2,  # Service Fee
    "T": NUM2,  # Fuel
    "U": NUM2,  # Surcharge
    "V": NUM2,  # Sub-Total
    "W": DEC2,  # R/kg
}


def _rgb(c):
    return c.rgb[-6:] if c is not None and isinstance(c.rgb, str) else None


def _stored_width(set_width: float) -> float:
    """xlsxwriter's width → pixels → stored-width conversion (Calibri 11:
    max digit width 7px, 5px padding), so the expectation is stated in the
    units the writer was given."""
    pixels = int(set_width * 7 + 0.5) + 5
    return int(pixels * 256 / 7) / 256  # stored to 1/256 of a character


def _widths_by_column(ws, n=23):
    """openpyxl collapses a run of equal-width columns into one range keyed
    by its first letter; expand it back to one width per column."""
    out = [None] * n
    for dim in ws.column_dimensions.values():
        for idx in range(dim.min, dim.max + 1):
            if idx <= n:
                out[idx - 1] = dim.width
    return out


@pytest.fixture(scope="module")
def tab(tmp_path_factory):
    out = tmp_path_factory.mktemp("fmt") / "tab.xlsx"
    wb = xlsxwriter.Workbook(str(out))
    st = Styles(wb)
    ix, day_rows, _ = build_billing_detail.load_day(None, DAY, (INV_HEADERS, ROWS))
    build_billing_detail.write_billing_tab(wb, st, DAY, ix, day_rows)
    wb.close()
    return openpyxl.load_workbook(out)["Billing 31 Jul"]


class TestSheetProperties:
    def test_tab_colour_is_navy(self, tab):
        assert _rgb(tab.sheet_properties.tabColor) == NAVY.lstrip("#").upper()

    def test_headings_held_at_b8(self, tab):
        assert tab.freeze_panes == "B8"

    def test_gridlines_hidden(self, tab):
        assert tab.sheet_view.showGridLines is False

    def test_the_23_column_widths(self, tab):
        assert _widths_by_column(tab) == pytest.approx(
            [_stored_width(w) for w in WIDTHS]
        )

    def test_title_block_and_heading_row_heights(self, tab):
        assert [tab.row_dimensions[r].height for r in (2, 3, 4, 5)] == [24, 30, 16, 4]
        assert tab.row_dimensions[7].height == pytest.approx(21.9)


class TestHeadingRow:
    def test_navy_fill_white_bold_9_wrapped_centred(self, tab):
        for col in range(2, 24):
            c = tab.cell(7, col)
            assert _rgb(c.fill.start_color) == "05003C", c.value
            assert _rgb(c.font.color) == "FFFFFF", c.value
            assert c.font.bold and c.font.size == 9, c.value
            assert c.alignment.wrap_text and c.alignment.horizontal == "center", c.value

    def test_the_22_headings_in_order(self, tab):
        assert [tab.cell(7, c).value for c in range(2, 24)] == [
            "Waybill", "Date", "Invoice #", "Account", "Customer", "Consignor",
            "Consignee", "Orig", "Dest", "Destination", "Cust Ref", "Service",
            "Pcs", "Actual Mass", "Chrge Mass", "Basic", "Outly", "Service Fee",
            "Fuel", "Surcharge", "Sub-Total", "R/kg",
        ]  # fmt: skip


class TestDataRows:
    def test_number_formats(self, tab):
        for letter, fmt in NUMBER_FORMATS.items():
            assert tab[f"{letter}8"].number_format == fmt, letter
        for letter in "BCDEFGHIJKLM":
            assert tab[f"{letter}8"].number_format == "General", letter

    def test_source_figures_blue_derived_black(self, tab):
        for letter in "NOPQRSTUV":  # Pcs .. Sub-Total
            assert _rgb(tab[f"{letter}8"].font.color) == "0000FF", letter
        assert tab["W8"].font.color is None or _rgb(tab["W8"].font.color) in (
            None,
            "000000",
        )

    def test_size_9_and_alternate_banding(self, tab):
        """Rows 10-11 are CUSTOMER 1's two lines (A0's single line sits on
        row 8 with its subtotal on 9); odd 0-indexed rows are white, even ALT."""
        assert tab["B8"].font.size == 9 and tab["B10"].font.size == 9
        assert tab["F10"].value == "CUSTOMER 1" and tab["F11"].value == "CUSTOMER 1"
        assert _rgb(tab["B10"].fill.start_color) == "FFFFFF"
        assert _rgb(tab["B11"].fill.start_color) == ALT.lstrip("#").upper()

    def test_every_populated_cell_is_boxed(self, tab):
        for r in (7, 8, tab.max_row):
            for c in range(2, 24):
                assert tab.cell(r, c).border.left.style == "thin", f"r{r}c{c}"


class TestTotalsRows:
    def _row(self, tab, label_fragment):
        return next(
            r
            for r in range(8, tab.max_row + 1)
            if label_fragment in str(tab.cell(r, 6).value or "")
        )

    def test_customer_subtotal_is_the_yellow_band(self, tab):
        r = self._row(tab, "subtotal")
        for c in range(2, 24):
            cell = tab.cell(r, c)
            assert _rgb(cell.fill.start_color) == YELLOW.lstrip("#").upper(), c
            assert (
                cell.font.bold and _rgb(cell.font.color) == NAVY.lstrip("#").upper()
            ), c
        assert (
            tab.cell(r, 22).number_format == NUM2
            and tab.cell(r, 23).number_format == DEC2
        )

    def test_grand_total_is_the_orange_band_on_the_last_row(self, tab):
        r = self._row(tab, "GRAND TOTAL")
        assert r == tab.max_row
        for c in range(2, 24):
            cell = tab.cell(r, c)
            assert _rgb(cell.fill.start_color) == ORANGE.lstrip("#").upper(), c
            assert (
                cell.font.bold and _rgb(cell.font.color) == NAVY.lstrip("#").upper()
            ), c
        for letter, fmt in NUMBER_FORMATS.items():
            assert tab[f"{letter}{r}"].number_format == fmt, letter
