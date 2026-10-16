"""Tab 1 of the Revenue Dashboard must lay out for every month of the year.

The month table on "YTD Revenue vs PY" grows by one row per completed month,
and the momentum block below it moves down with it. The three footnotes under
the momentum block were pinned to rows 33-35 — where they happened to fall
with six completed months. With seven (October 2026, Mar-Sep complete) the
eighth momentum row reached row 33 and xlsxwriter refused the overlapping
merge, so the 16:30 pack failed on 14 and 15 Oct 2026, the first runs with a
billing day in October.

That is a layout that only breaks in a month nobody built, so this builds
tab 1 for the shapes the year can take — with no export behind it, just a
model that answers the three totals the tab asks for — and checks the
footnotes land directly under the momentum block every time.
"""

import sys
from datetime import date
from pathlib import Path

import pytest
import xlsxwriter
from openpyxl import load_workbook

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO / "revenue_reports"))

from build_dashboard import build_tab1, period_params
from style import Styles


class _Pool:
    """Stands in for build_dashboard.Pool: tab 1 only asks for window totals."""

    def __init__(self, per_day: float, kg_per_day: float):
        self.per_day, self.kg_per_day = per_day, kg_per_day

    def total(self, a: date, b: date):
        n = (b - a).days + 1
        return self.per_day * n, self.kg_per_day * n


class _Credits:
    def __init__(self, per_day: float):
        self.per_day = per_day

    def total(self, a: date, b: date):
        return self.per_day * ((b - a).days + 1)


class _Model:
    inv = _Pool(100_000.0, 1_000.0)
    py = _Pool(90_000.0, 950.0)
    credits = _Credits(2_000.0)


def build_tab1_for(mtd_max: date, out: Path) -> Path:
    path = out / f"tab1-{mtd_max.isoformat()}.xlsx"
    wb = xlsxwriter.Workbook(str(path))
    try:
        build_tab1(wb, Styles(wb), _Model(), period_params(mtd_max))
    finally:
        wb.close()  # OverlappingRange is raised by merge_range, before close
    return path


def merged_text(path: Path):
    """[(min_row, max_row, min_col, max_col, text)] for every merge on tab 1."""
    ws = load_workbook(path).worksheets[0]
    out = []
    for r in ws.merged_cells.ranges:
        out.append(
            (
                r.min_row,
                r.max_row,
                r.min_col,
                r.max_col,
                ws.cell(r.min_row, r.min_col).value,
            )
        )
    return sorted(out)


# One date per shape the year takes. Each is a plausible "latest invoiced
# day" and the number of completed months is asserted, so a change to the
# month-selection rule shows up here as the wrong shape, not a silent pass.
SHAPES = [
    pytest.param(date(2026, 9, 15), 6, id="sep-mtd-6-complete-the-shape-that-worked"),
    pytest.param(date(2026, 9, 30), 7, id="sep-month-end-7-complete-no-part-month"),
    pytest.param(date(2026, 10, 15), 7, id="oct-mtd-7-complete-the-shape-that-failed"),
    pytest.param(date(2027, 2, 15), 11, id="feb-mtd-11-complete-last-month-of-fy"),
    pytest.param(date(2027, 2, 28), 12, id="feb-month-end-12-complete-full-year"),
]


class TestFootnotesFollowTheMomentumBlock:
    @pytest.mark.parametrize("mtd_max, n_complete", SHAPES)
    def test_builds_without_overlapping_merges(self, tmp_path, mtd_max, n_complete):
        assert len(period_params(mtd_max)["complete"]) == n_complete
        build_tab1_for(mtd_max, tmp_path)  # raises OverlappingRange if it regresses

    @pytest.mark.parametrize("mtd_max, n_complete", SHAPES)
    def test_footnotes_sit_directly_under_the_last_momentum_row(
        self, tmp_path, mtd_max, n_complete
    ):
        merges = merged_text(build_tab1_for(mtd_max, tmp_path))
        p = period_params(mtd_max)
        last_momentum = [
            m
            for m in merges
            if str(m[4]).startswith("Projected") and "growth vs PY" in m[4]
        ]
        assert len(last_momentum) == 1
        last_row = last_momentum[0][1]

        footnotes = [m for m in merges if m[0] > last_row]
        assert [(m[0], m[2], m[3]) for m in footnotes] == [
            (last_row + 1, 2, 10),  # B:J, the row after the momentum block
            (last_row + 2, 2, 10),
            (last_row + 3, 2, 10),
        ], "three footnotes, B:J, starting on the row after the momentum block"

        # The block above moves one row per completed month, so the footnotes
        # must too: 33 with six completed months is the layout that shipped
        # through September 2026.
        expected_first = (
            14 + n_complete + (2 if p["cur"] not in p["complete"] else 0) + 11
        )
        assert footnotes[0][0] == expected_first
        if n_complete == 6:
            assert footnotes[0][0] == 33

    @pytest.mark.parametrize("mtd_max, n_complete", SHAPES)
    def test_first_footnote_names_the_current_month(
        self, tmp_path, mtd_max, n_complete
    ):
        merges = merged_text(build_tab1_for(mtd_max, tmp_path))
        p = period_params(mtd_max)
        first = [m[4] for m in merges if str(m[4]).startswith("¹")]
        assert len(first) == 1
        text = first[0]
        month = mtd_max.strftime("%B")
        assert text.startswith(
            f"¹ {p['fyl']} {month} = {p['elapsed']} invoiced trading days"
        )
        assert f"{p['ly']} {month} shown like-for-like" in text
        assert f"full {p['ly']} {month} net = " in text
        if month != "July":
            assert "July" not in text
